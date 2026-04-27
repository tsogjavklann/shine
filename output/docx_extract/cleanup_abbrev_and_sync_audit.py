from datetime import datetime
from math import exp
from pathlib import Path
import csv
import shutil
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


DOCX = Path("bagiin_ner_v29.docx")
BACKUP_DIR = Path("output/docx_backups")
REPORT_DIR = Path("output/reports")
BACKUP_DIR.mkdir(parents=True, exist_ok=True)
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_run_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_para_text(paragraph, text, *, size=12, bold=False, italic=False, align=None):
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    if align is not None:
        paragraph.alignment = align


def set_cell(cell, text, *, header=False):
    p = cell.paragraphs[0]
    clear_runs(p)
    r = p.add_run(text)
    set_run_font(r, size=9, bold=header)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.LEFT


def delete_row(table, row):
    tbl = table._tbl
    tbl.remove(row._tr)


def row_text(row):
    return [" ".join(c.text.split()) for c in row.cells]


def read_csv_one(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


backup = BACKUP_DIR / f"bagiin_ner_v29_before_abbrev_sync_cleanup_{datetime.now():%Y%m%d_%H%M%S}.docx"
shutil.copy2(DOCX, backup)

doc = Document(DOCX)

# Title: avoid the foreign two-letter "IV" in the main title.
for p in doc.paragraphs:
    if "IV-THRESHOLD РЕГРЕССИЙН ШИНЖИЛГЭЭ" in p.text:
        set_para_text(
            p,
            "МОНГОЛ УЛСАД БОЛОВСРОЛЫН БОДИТ ӨГӨӨЖИЙН БОСГОТОЙ ҮНЭЛГЭЭ: "
            "ӨРХИЙН НИЙГЭМ ЭДИЙН ЗАСГИЙН СУДАЛГААНД СУУРИЛСАН "
            "ХЭРЭГСЭЛ ХУВЬСАГЧТАЙ БОСГО УТГАТ РЕГРЕССИЙН ШИНЖИЛГЭЭ",
            size=12,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        break

# Remove the newly introduced two-letter abbreviation "ТН".
abbr_table = doc.tables[1]
for row in list(abbr_table.rows):
    if row_text(row)[0] == "ТН":
        delete_row(abbr_table, row)

# Make the variable definition table read like prose, not code.
var_table = doc.tables[3]
symbol_replacements = {
    "Боловсролын жил": "боловсролын жил",
    "Эцэг, эхийн боловсролын дундаж": "эцэг, эхийн боловсролын дундаж",
    "Сурагч-багшийн харьцаа 17–18 нас": "босго хувьсагч",
    "Хяналтын хувьсагчид": "хяналтын хувьсагчид",
    "Тогтмол нөлөө": "тогтмол нөлөө",
    "Түүврийн жин": "түүврийн жин",
}
for row in var_table.rows:
    cells = row_text(row)
    if cells[0] in symbol_replacements:
        set_cell(row.cells[1], symbol_replacements[cells[0]])

# Replace code-ish wording in result tables.
first_stage_table = doc.tables[5]
for row in first_stage_table.rows:
    cells = row_text(row)
    if cells[0] == "Сул IV шалгуур":
        set_cell(row.cells[0], "Сул хэрэгслийн шалгуур")
        set_cell(row.cells[1], "F ≥ 10")
        set_cell(row.cells[5], "Сул хэрэгсэл гэж үзэхгүй")

boot_table = doc.tables[7]
for row in boot_table.rows:
    cells = row_text(row)
    if cells[0] == "Bootstrap p-утга":
        set_cell(row.cells[0], "Бүүтстрап p-утга")
    if len(cells) > 3 and cells[3] == "B=399; B=1999 шаардлагатай":
        set_cell(row.cells[3], "399 давталт; 1,999 давталт шаардлагатай")

summary_table = doc.tables[8]
for row in summary_table.rows:
    cells = row_text(row)
    if cells[0] == "Регимийн зөрүү":
        set_cell(row.cells[3], "Wald p=0.102; бүүтстрап p=0.0451")
        set_cell(row.cells[5], "399 давталт тул урьдчилсан")
    if cells[0] == "Бүүтстрап":
        set_cell(row.cells[3], "алдаагүй")
    if cells[0] == "Эхний шатны F":
        set_cell(row.cells[5], "Сул хэрэгсэл биш")
    if cells[0] == "Тайлбар":
        set_cell(row.cells[5], "Шууд шалтгаан гэж тайлбарлахгүй")

# Remove short code symbols from prose and keep the trimming rule synced to R/28.
for p in doc.paragraphs:
    text = " ".join(p.text.split())
    if text.startswith("Энд PE нь эцэг, эхийн боловсролын дундаж"):
        set_para_text(
            p,
            "Энд эцэг, эхийн боловсролын дундажийг тухайн ажиллагчийн эцэг болон эхийн "
            "ажиглагдсан боловсролын жилүүдийн арифметик дундажаар тооцсон. Хэрэв эцэг "
            "эсвэл эхийн аль нэгний боловсролын мэдээлэл ажиглагдсан бол тухайн ажиглагдсан "
            "мэдээлэлд тулгуурлаж, хоёулаа байхгүй тохиолдолд хувьсагчийг байгуулаагүй. "
            "Энэ хувьсагч боловсролын сонголтод урьдчилан нөлөөлөх гэр бүлийн суурь нөхцөлийг "
            "төлөөлөх боловч цалинд гэр бүлийн сүлжээ, хүмүүжил, ажиглагдахгүй чадвараар "
            "дамжин шууд нөлөөлөх боломжийг бүрэн үгүйсгэхгүй тул тайлбар нь болгоомжтой байна.",
        )
    if "Хайлтыг босго хувьсагчийн 15–85 хувийн мужид" in text:
        set_para_text(
            p,
            "ХХБР үнэлгээний дараагийн алхам нь 17–18 насны сурагч-багшийн дундаж харьцааны "
            "боломжит босго утгуудыг сүлжээн хайлтаар шалгах явдал юм. Тогтмол нөлөөг "
            "үлдэгдэлжүүлсний дараа боломжит босго бүр дээр хоёр регимийн хэрэгсэл хувьсагчтай "
            "налууг үнэлж, нэгтгэн агшаасан алдааны квадратын нийлбэрийг харьцуулсан. "
            "Хайлтыг босго хувьсагчийн 10–90 хувийн мужид хязгаарласнаар хоёр регимийн аль "
            "нэгэнд хэт цөөн ажиглалт үлдэхээс сэргийлсэн. Нийт 278 боломжит босго шалгагдаж, "
            "зорилгын функцийн хамгийн бага утга 19.53 сурагч/багш дээр тогтов.",
        )

doc.save(DOCX)

# ---------------------------------------------------------------------------
# Code-result synchronization audit
# ---------------------------------------------------------------------------
doc2 = Document(DOCX)
all_text = "\n".join(" ".join(p.text.split()) for p in doc2.paragraphs)
table_text = "\n".join(
    " | ".join(" ".join(c.text.split()) for c in row.cells)
    for tbl in doc2.tables
    for row in tbl.rows
)
full_text = all_text + "\n" + table_text

baseline = read_csv_one("output/tables/T14a_student_teacher_avg_17_18_baseline_ols_2sls.csv")
fs = read_csv_one("output/tables/T14a_student_teacher_avg_17_18_parent_iv_first_stage.csv")[0]
diag = read_csv_one("output/tables/T14a_student_teacher_avg_17_18_threshold_sample_diagnostics.csv")[0]
gamma = read_csv_one("output/tables/T14c_student_teacher_avg_17_18_ch_gamma_hat.csv")[0]
gmm = read_csv_one("output/tables/T14d_student_teacher_avg_17_18_ch_gmm_final_results.csv")
wald = read_csv_one("output/tables/T14d_student_teacher_avg_17_18_ch_gmm_wald_test.csv")[0]
boot = read_csv_one("output/tables/T14e_student_teacher_avg_17_18_ch_bootstrap_inference.csv")[0]

ols = next(r for r in baseline if r["model"] == "OLS baseline")
iv = next(r for r in baseline if r["model"] == "2SLS parent_educ_mean IV")
low = next(r for r in gmm if r["term"] == "educ_low")
high = next(r for r in gmm if r["term"] == "educ_high")


def f(x):
    return float(x)


expected = {
    "OLS N": f'{int(float(ols["N"])):,}',
    "OLS beta": f'{f(ols["estimate"]):.4f}',
    "OLS se": f'{f(ols["se"]):.4f}',
    "OLS return": f'{100 * (exp(f(ols["estimate"])) - 1):.1f}%',
    "2SLS N": f'{int(float(iv["N"])):,}',
    "2SLS beta": f'{f(iv["estimate"]):.4f}',
    "2SLS se": f'{f(iv["se"]):.4f}',
    "2SLS return": f'{100 * (exp(f(iv["estimate"])) - 1):.1f}%',
    "first stage coef": f'{f(fs["estimate"]):.4f}',
    "first stage se": f'{f(fs["se"]):.4f}',
    "first stage t": f'{f(fs["t_stat"]):.2f}',
    "first stage F": f'{f(fs["first_stage_F"]):.2f}',
    "gamma": f'{f(gamma["gamma_hat"]):.2f}',
    "N low": f'{int(float(gamma["N_low"])):,}',
    "N high": f'{int(float(gamma["N_high"])):,}',
    "beta low": f'{f(low["estimate"]):.4f}',
    "se low": f'{f(low["se"]):.4f}',
    "return low": f'{100 * (exp(f(low["estimate"])) - 1):.1f}%',
    "beta high": f'{f(high["estimate"]):.4f}',
    "se high": f'{f(high["se"]):.4f}',
    "return high": f'{100 * (exp(f(high["estimate"])) - 1):.1f}%',
    "beta diff": f'{f(boot["beta_diff_observed"]):.4f}',
    "beta diff pp": f'{100 * f(boot["beta_diff_observed"]):.1f}',
    "wald p": f'{f(wald["p_value"]):.3f}',
    "bootstrap p": f'{f(boot["bootstrap_p_value"]):.4f}',
    "boot B": str(int(float(boot["B_requested"]))),
    "boot success": str(int(float(boot["n_success"]))),
    "boot failed": str(int(float(boot["n_failed"]))),
    "gamma q025": f'{f(boot["gamma_q025"]):.2f}',
    "gamma q50": f'{f(boot["gamma_q50"]):.2f}',
    "gamma q975": f'{f(boot["gamma_q975"]):.2f}',
    "beta diff ci low": f'{f(boot["beta_diff_q025"]):.4f}',
    "beta diff ci high": f'{f(boot["beta_diff_q975"]):.4f}',
    "sample clusters": str(int(float(diag["n_birth_aimag_clusters"]))),
    "sample cohorts": str(int(float(diag["n_birth_cohort_groups"]))),
    "sample waves": str(int(float(diag["n_waves"]))),
    "age min": str(int(float(diag["age_min"]))),
    "age max": str(int(float(diag["age_max"]))),
    "birth year min": str(int(float(diag["birth_year_min"]))),
    "birth year max": str(int(float(diag["birth_year_max"]))),
    "year at 17/18 min": str(int(float(diag["year_at_17_18_min"]))),
    "year at 17/18 max": str(int(float(diag["year_at_17_18_max"]))),
    "q min": f'{f(diag["student_teacher_ratio_avg_17_18_min"]):.2f}',
    "q median": f'{f(diag["student_teacher_ratio_avg_17_18_p50"]):.2f}',
    "q max": f'{f(diag["student_teacher_ratio_avg_17_18_max"]):.2f}',
    "corr educ": f'{f(diag["corr_student_teacher_ratio_avg_17_18_educ_years"]):.3f}'.replace("-", "−"),
    "corr parent": f'{f(diag["corr_student_teacher_ratio_avg_17_18_parent_educ_mean"]):.3f}',
    "corr lwage": f'{f(diag["corr_student_teacher_ratio_avg_17_18_lwage"]):.3f}',
}

audit_rows = []
for name, value in expected.items():
    audit_rows.append({
        "item": name,
        "expected_from_code": value,
        "found_in_docx": str(value in full_text),
    })

bad_terms = ["ТН", " PE ", "B=399", "B=1999", "0 failed", "Weak-IV flag", "IV-THRESHOLD РЕГРЕССИЙН ШИНЖИЛГЭЭ", "15–85", "15-85"]
for term in bad_terms:
    audit_rows.append({
        "item": f"forbidden_or_codeish_term::{term}",
        "expected_from_code": "absent",
        "found_in_docx": str(term in full_text),
    })

audit_csv = REPORT_DIR / "docx_code_sync_audit.csv"
with audit_csv.open("w", newline="", encoding="utf-8-sig") as fcsv:
    writer = csv.DictWriter(fcsv, fieldnames=["item", "expected_from_code", "found_in_docx"])
    writer.writeheader()
    writer.writerows(audit_rows)

missing = [r for r in audit_rows if not r["item"].startswith("forbidden_or_codeish_term::") and r["found_in_docx"] != "True"]
forbidden_present = [r for r in audit_rows if r["item"].startswith("forbidden_or_codeish_term::") and r["found_in_docx"] == "True"]

audit_md = REPORT_DIR / "docx_code_sync_audit.md"
audit_md.write_text(
    "\n".join([
        "# DOCX-Code Synchronization Audit",
        "",
        f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}",
        f"DOCX: {DOCX}",
        f"Backup: {backup}",
        "",
        f"Checked items: {len(expected)}",
        f"Missing expected rounded values in docx: {len(missing)}",
        f"Forbidden/code-like terms still present: {len(forbidden_present)}",
        "",
        "## Missing Expected Values",
        *(f"- {r['item']}: {r['expected_from_code']}" for r in missing),
        "",
        "## Forbidden Terms Present",
        *(f"- {r['item']}" for r in forbidden_present),
        "",
        "## Note",
        "The audit checks the rounded values used in the report against the T14 student_teacher_ratio_avg_17_18 output tables.",
    ]),
    encoding="utf-8",
)

with ZipFile(DOCX) as z:
    zip_ok = z.testzip() is None

print(f"saved={DOCX}")
print(f"backup={backup}")
print(f"audit_csv={audit_csv}")
print(f"audit_md={audit_md}")
print(f"zip_ok={zip_ok}")
print(f"missing_expected_values={len(missing)}")
print(f"forbidden_terms_present={len(forbidden_present)}")
