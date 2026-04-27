from __future__ import annotations

import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOCX_PATH = Path("bagiin_ner_v29.docx")
BACKUP_DIR = Path("output/docx_backups")
FALLBACK_PATH = Path("output/bagiin_ner_v29_sudlagdsan_updated.docx")

LITERATURE_ITEMS = [
    "1.1. Боловсролын өгөөж ба эндоген хазайлт",
    "Боловсролын цалингийн өгөөжийг тайлбарлах суурь онол нь хүмүүн капиталын онол юм. Schultz (1961), Becker (1964) нар боловсролыг хувь хүний бүтээмжийг нэмэгдүүлэх хөрөнгө оруулалт гэж үзсэн бол Mincer (1974) боловсролын жил, хөдөлмөрийн туршлагыг цалингийн түвшинтэй холбосон эмпирик тэгшитгэлийг боловсруулсан. Энэхүү аргачлал нь олон улсын хөдөлмөрийн эдийн засгийн судалгаанд өргөн хэрэглэгдэж, Psacharopoulos, Patrinos (2018)-ийн мета-шинжилгээгээр боловсролын нэмэлт нэг жилийн дундаж өгөөж 9.0 хувь орчим, дээд боловсролын өгөөж үүнээс өндөр байдгийг харуулсан байдаг.",
    "Гэвч Минсерийн тэгшитгэлийг ЭХБК аргаар шууд үнэлэхэд боловсролын сонголт санамсаргүй бус байдгаас хазайлт үүсдэг. Нэгдүгээрт, ажиглагдахгүй ур чадвар, суралцах чадвар, хичээл зүтгэл зэрэг шинжүүд боловсролын жил болон цалинд зэрэг нөлөөлөх тул боловсролын коэффициент дээш чиглэсэн хазайлттай байж болно. Хоёрдугаарт, боловсролын жил буруу хэмжигдэхэд сулруулагч хазайлт үүсэх боломжтой. Card (1999) эдгээр эсрэг чиглэлтэй хазайлтыг нэгэн зэрэг авч үзэх шаардлагатайг онцолсон бөгөөд иймээс боловсролын өгөөжийг зөвхөн ЭХБК үнэлгээнд тулгуурлан тайлбарлах нь хангалтгүй юм.",
    "1.2. Хэрэгсэл хувьсагч ба гэр бүлийн суурь нөхцөл",
    "Эндоген хазайлтыг бууруулах гол арга нь боловсролын жилд нөлөөлөх боловч цалинд боловсролоор дамжин нөлөөлөх хэрэгсэл хувьсагч ашиглах явдал юм. Angrist, Krueger (1991) заавал суралцах насны хууль, Card (1993) коллежийн газарзүйн ойртоц, Duflo (2001) сургуулийн барилгын хөтөлбөрийг ашиглан боловсролын шалтгаант өгөөжийг үнэлсэн нь энэ чиглэлийн үндсэн жишээ болсон. Эдгээр судалгаанд ХШХБК үнэлгээ ЭХБК-ийн үнэлгээнээс ялгаатай гардаг нь боловсролын сонголтын эндоген шинжийг эмпирикээр засах шаардлагатайг харуулдаг.",
    "Энэ судалгаанд боловсролын жилийн хэрэгсэл хувьсагчаар эцэг, эхийн боловсролын дундажийг ашиглаж байна. Эцэг эхийн боловсрол нь хүүхдийн суралцах орчин, боловсрол үргэлжлүүлэх хандлага, мэдээлэл авах боломж, өрхийн боловсролд хандах үнэлэмжээр дамжин хүүхдийн боловсролын жилтэй хүчтэй холбоотой байх онолын үндэслэлтэй. Харин уг хувьсагч цалинд гэр бүлийн нийгмийн сүлжээ, хүмүүжил, ажиглагдахгүй чадвар зэрэг сувгаар шууд нөлөөлөх боломжийг бүрэн үгүйсгэх боломжгүй тул үр дүнг Conley, Hansen, Rossi (2012)-ийн онцолдог болгоомжтой хэрэгсэл хувьсагчийн тайлбарын хүрээнд авч үзнэ. Иймээс энэхүү ажилд эцэг, эхийн боловсролын дундажийг хүчтэй боловч бүрэн төгс бус хэрэгсэл хувьсагч гэж ил тод авч үзсэн.",
    "1.3. Сургуулийн орчны босго ба ХХБР арга",
    "Боловсролын өгөөж нь зөвхөн хувь хүний боловсролын жилийн тооноос бус тухайн боловсрол хуримтлагдсан сургуулийн орчноос хамаарч өөр байж болно. Сургуулийн нөөц, багшийн ачаалал, ангийн нягтаршил нь суралцах үйл явцад нөлөөлөх боломжтойг Duflo (2001)-ийн сургуулийн нийлүүлэлтийн судалгаа, Angrist, Lavy (1999) болон Krueger (1999)-ийн ангийн хэмжээний судалгаанууд харуулсан. Hanushek (2003) сургуулийн орцын нөлөө улс орон, байгууллагын чанар, хэрэгжилтийн нөхцөлөөс хамааран тогтвортой бус байж болохыг сануулсан нь сурагч-багшийн харьцааг шууд шалтгаант нөлөө бус, сургуулийн орчны ялгааг илэрхийлэх босго хувьсагч болгон болгоомжтой ашиглах үндэслэл болно.",
    "Босготой регрессийн арга нь өгөгдлийг тодорхой босго хувьсагчийн утгаар хоёр болон түүнээс олон регимд хувааж, регим бүрт ялгаатай параметр үнэлэх боломж олгодог. Hansen (2000) экзоген тайлбарлагчтай нөхцөлд босго утгыг өгөгдлөөс хайх, босго байхгүй гэсэн таамаглалыг дахин түүвэрлэлтийн аргаар шалгах аргачлалыг боловсруулсан. Энэ арга нь шугаман нэг налуугаар тайлбарлагдахгүй бүтцийн ялгааг илрүүлэхэд ач холбогдолтой боловч боловсрол зэрэг эндоген тайлбарлагчтай үед дангаараа хангалтгүй юм.",
    "Caner, Hansen (2004) нарын хэрэгсэл хувьсагчтай босго утгат регресс нь боловсролын эндоген шинж ба өгөөжийн регимийн ялгааг нэгэн зэрэг авч үзэх боломж олгодог. Энэ судалгаанд боловсролын жил нь эндоген үндсэн тайлбарлагч, эцэг, эхийн боловсролын дундаж нь хэрэгсэл хувьсагч, харин 17–18 насны сурагч-багшийн дундаж харьцаа нь босго хувьсагчийн үүрэгтэй. Ийм бүтэц нь боловсролын өгөөж сургуулийн орчны ялгаагаар өөрчлөгдөж байгаа эсэхийг шалгах боломж олгох бөгөөд босго хувьсагчийг цалингийн шууд шалтгаан гэж бус, өгөөжийн ялгаатай регимийг ялгах шалгуур гэж тайлбарлана.",
    "1.4. Монголын хүрээн дэх өмнөх судалгаа ба орон зай",
    "Монгол Улсад боловсролын өгөөжийг үнэлсэн эмпирик судалгаа харьцангуй цөөн боловч бодлогын ач холбогдол өндөртэй байна. Pastore (2010) 2007–2008 оны өрхийн түүвэрт тулгуурлан залуучуудын боловсролын өгөөжийг 7–9 хувь орчим гэж үнэлсэн. Гэвч уг ажил нь голлон ЭХБК хэлбэрийн үнэлгээнд суурилсан бөгөөд боловсролын эндоген сонголт, сургуулийн орчны босго шинж, регим хоорондын ялгаатай өгөөжийг хамтад нь авч үзээгүй.",
    "Ай Ар Ай Эм ХХК (2015)-ийн хөдөлмөрийн зах зээлийн шинжилгээ Ажиллах хүчний судалгааны өгөгдөлд тулгуурлан дээд боловсролтой иргэдийн цалингийн давуу талыг тооцсон. Мөн Montenegro, Patrinos (2014)-ийн олон улсын харьцуулалт Монгол Улсын боловсролын өгөөжийг бусад улс оронтой жиших суурь өгдөг. Гэсэн хэдий ч эдгээр ажилд хэрэгсэл хувьсагч, босго утгат регрессийг нэгтгэсэн үнэлгээ хийгдээгүй тул боловсролын шалтгаант өгөөж сургуулийн орчны өөр өөр регимд хэрхэн ялгаатай болох асуудал нээлттэй хэвээр байна.",
    "Иймээс Монголын хөдөлмөрийн эдийн засгийн уран зохиолд өрхийн түвшний өгөгдөл дээр боловсролын эндоген шинжийг эцэг, эхийн боловсролын дундаж хэрэгсэл хувьсагчаар засаж, 17–18 насны сурагч-багшийн дундаж харьцаагаар тодорхойлогдох босго регимүүдийн хооронд боловсролын өгөөжийг харьцуулсан судалгаа дутагдаж байна. Энэхүү судалгаа ӨНЭЗС-ийн 2020, 2021, 2022, 2024 оны давалгааг ашиглан Caner, Hansen (2004)-ийн санаанд тулгуурласан ХХБР аргачлалыг Монголын өгөгдөлд хэрэглэж буйгаараа уг орон зайг нөхөхийг зорьж байна.",
]

REFERENCE_ITEMS = [
    "Angrist, J. D., & Krueger, A. B. (1991). Does compulsory school attendance affect schooling and earnings? Quarterly Journal of Economics, 106(4), 979–1014.",
    "Angrist, J. D., & Lavy, V. (1999). Using Maimonides' rule to estimate the effect of class size on scholastic achievement. Quarterly Journal of Economics, 114(2), 533–575.",
    "Hanushek, E. A. (2003). The failure of input-based schooling policies. Economic Journal, 113(485), F64–F98.",
    "Krueger, A. B. (1999). Experimental estimates of education production functions. Quarterly Journal of Economics, 114(2), 497–532.",
]


def normalized(text: str) -> str:
    return " ".join(text.split()).upper()


def insert_paragraph_before(paragraph: Paragraph, text: str, style_name: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style_name:
        new_para.style = style_name
    new_para.text = text
    return new_para


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def ensure_reference_order(doc: Document) -> None:
    first_ref = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("Becker, G. S.")
    )
    style_name = first_ref.style.name

    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip() in REFERENCE_ITEMS:
            remove_paragraph(paragraph)

    first_ref = next(
        p for p in doc.paragraphs
        if p.text.strip().startswith("Becker, G. S.")
    )
    for ref in REFERENCE_ITEMS:
        insert_paragraph_before(first_ref, ref, style_name)


def main() -> None:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"bagiin_ner_v29_before_sudlagdsan_baidal_{stamp}.docx"
    shutil.copy2(DOCX_PATH, backup)

    doc = Document(DOCX_PATH)
    start = next(
        i for i, p in enumerate(doc.paragraphs)
        if "СУДЛАГДСАН БАЙДАЛ" in normalized(p.text)
    )
    end = next(
        i for i in range(start + 1, len(doc.paragraphs))
        if normalized(doc.paragraphs[i].text).startswith("II БҮЛЭГ")
    )
    slots = [i for i in range(start + 1, end) if doc.paragraphs[i].text.strip()]
    if len(slots) != len(LITERATURE_ITEMS):
        raise RuntimeError(f"Expected {len(LITERATURE_ITEMS)} literature slots, found {len(slots)}")

    for idx, text in zip(slots, LITERATURE_ITEMS):
        doc.paragraphs[idx].text = text

    ensure_reference_order(doc)

    saved_path = DOCX_PATH
    try:
        doc.save(DOCX_PATH)
    except PermissionError:
        FALLBACK_PATH.parent.mkdir(parents=True, exist_ok=True)
        doc.save(FALLBACK_PATH)
        saved_path = FALLBACK_PATH

    with zipfile.ZipFile(saved_path) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Invalid docx zip member: {bad}")
    print(f"Updated literature review in {saved_path}")
    print(f"Backup: {backup}")
    print(f"Literature paragraphs replaced: {len(LITERATURE_ITEMS)}")
    if saved_path != DOCX_PATH:
        print("Original was locked; fallback copy was written.")


if __name__ == "__main__":
    main()
