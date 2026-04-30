from __future__ import annotations

import csv
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "bagiin_ner_v29.docx"
OUT = ROOT / "output" / "bagiin_ner_instrumental_forest_tailan.docx"
TABLE_DIR = ROOT / "output" / "tables"
FIG_DIR = ROOT / "output" / "figures"


def read_csv_dict(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


summary = {r["item"]: float(r["value"]) for r in read_csv_dict(TABLE_DIR / "T16_instrumental_forest_summary.csv")}
quartiles = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_tau_quartiles.csv")
regimes = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_threshold_regimes.csv")
importance = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_variable_importance.csv")


def pct(x: float) -> str:
    return f"{float(x):.2f}"


def num(x: float, digits: int = 2) -> str:
    return f"{float(x):.{digits}f}"


def clear_document(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def set_page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(1.2)
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.text = ""
        add_page_number(p)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in [("Title", 20), ("Heading 1", 14), ("Heading 2", 12)]:
        if name in styles:
            st = styles[name]
            st.font.name = "Times New Roman"
            st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            st.font.size = Pt(size)
            st.font.bold = True
            st.paragraph_format.space_before = Pt(12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing = 1.5


def run_font(run, size: int | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_para(doc: Document, text: str = "", align=None, style=None, bold=False, italic=False, size=None, first_line=True):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    elif style is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if style is None and first_line and text:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    run_font(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    run_font(r, size=14 if level == 1 else 12, bold=True)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    run_font(r, size=10, bold=True, italic=True)


def add_source(doc: Document, text: str = "Эх сурвалж: Зохиогчийн тооцоолол.") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    run_font(r, size=9, italic=True)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade_cell(hdr[i], "D9EAF7")
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                run_font(r, size=10, bold=True)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    run_font(r, size=10)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def add_picture(doc: Document, image: Path, width_inches: float = 6.4) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image), width=Inches(width_inches))


def page_break(doc: Document) -> None:
    doc.add_page_break()


doc = Document(str(TEMPLATE))
clear_document(doc)
set_styles(doc)
set_page_setup(doc)

# Cover
add_para(doc, "САНХҮҮ ЭДИЙН ЗАСГИЙН ИХ СУРГУУЛЬ", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, first_line=False)
add_para(doc, "ЭКОНОМЕТРИКИЙН VIII ОЛИМПИАДЫН II ШАТ", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, first_line=False)
for _ in range(5):
    add_para(doc, "", first_line=False)
add_para(doc, "ЭРДЭМ ШИНЖИЛГЭЭНИЙ НЭМЭЛТ ТАЙЛАН", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, first_line=False)
add_para(doc, "Сэдэв: Боловсролын өгөөжийн ялгаатай байдлыг Instrumental Forest аргаар шалгах нь", WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=14, first_line=False)
add_para(doc, "Монгол Улсад боловсролын бодит өгөөжийн босготой үнэлгээ судалгаанд хавсаргах machine learning нэмэлт шинжилгээ", WD_ALIGN_PARAGRAPH.CENTER, size=12, first_line=False)
for _ in range(4):
    add_para(doc, "", first_line=False)
add_para(doc, "Боловсруулсан: А.Азжаргал, Ү.Номин, Ц.Цогжавхлан", WD_ALIGN_PARAGRAPH.CENTER, size=12, first_line=False)
add_para(doc, "Удирдагч: С.Өнөр", WD_ALIGN_PARAGRAPH.CENTER, size=12, first_line=False)
for _ in range(6):
    add_para(doc, "", first_line=False)
add_para(doc, "УЛААНБААТАР 2026", WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, first_line=False)
page_break(doc)

# Contents and lists
add_heading(doc, "АГУУЛГА", 1)
contents = [
    "ХУРААНГУЙ",
    "ОРШИЛ",
    "I БҮЛЭГ. СУДЛАГДСАН БАЙДАЛ",
    "II БҮЛЭГ. АРГА ЗҮЙ",
    "III БҮЛЭГ. ӨГӨГДӨЛ БА ХУВЬСАГЧИД",
    "IV БҮЛЭГ. ШИНЖИЛГЭЭНИЙ ҮР ДҮН",
    "V БҮЛЭГ. ДҮГНЭЛТ БА ТАЙЛАНД АШИГЛАХ ЗӨВЛӨМЖ",
    "НОМ ЗҮЙ",
    "ХАВСРАЛТ",
]
for item in contents:
    add_para(doc, item, first_line=False)

add_heading(doc, "ХҮСНЭГТЭН МЭДЭЭЛЛИЙН ЖАГСААЛТ", 1)
for item in [
    "Хүснэгт 1. Instrumental Forest шинжилгээний загварын бүтэц",
    "Хүснэгт 2. Local IV өгөөжийн үндсэн статистик",
    "Хүснэгт 3. Local IV өгөөжийн квартил бүлгүүдийн харьцуулалт",
    "Хүснэгт 4. STR босгоны регимээр харьцуулсан Instrumental Forest үр дүн",
    "Хүснэгт 5. Variable importance-ийн нэгтгэл",
]:
    add_para(doc, item, first_line=False)

add_heading(doc, "ЗУРГАН МЭДЭЭЛЛИЙН ЖАГСААЛТ", 1)
for item in [
    "Зураг 1. Instrumental Forest өгөөжийн тархалт",
    "Зураг 2. Local IV өгөөж ба сурагч-багшийн харьцааны хамаарал",
    "Зураг 3. Өгөөжийн ялгааг тайлбарлах хувьсагчдын ач холбогдол",
]:
    add_para(doc, item, first_line=False)

add_heading(doc, "ТОВЧИЛСОН ҮГС", 1)
add_table(
    doc,
    ["Товчлол", "Англи нэр", "Монгол тайлбар"],
    [
        ["IF", "Instrumental Forest", "Хэрэгсэл хувьсагчтай ой загвар"],
        ["IV", "Instrumental Variable", "Хэрэгсэл хувьсагч"],
        ["ХШХБК", "Two-Stage Least Squares", "Хоёр шатлалт хамгийн бага квадратын арга"],
        ["ХХБР", "IV Threshold Regression", "Хэрэгсэл хувьсагчтай босго утгат регресс"],
        ["STR", "Student-Teacher Ratio", "Нэг багшид ногдох сурагчийн тоо"],
    ],
    [2.5, 5.0, 8.0],
)
page_break(doc)

# Abstract
add_heading(doc, "ХУРААНГУЙ", 1)
add_para(
    doc,
    "Энэхүү нэмэлт тайлангаар Монгол Улсад боловсролын нэг жилийн цалингийн өгөөж хувь хүн, сургуулийн орчин, бүс нутгийн шинжүүдээр ялгаатай эсэхийг Instrumental Forest буюу хэрэгсэл хувьсагчтай ой загвар ашиглан эрэл хайгуулын байдлаар шалгав. Үндсэн судалгаанд боловсролын жил нь эндоген тайлбарлагч, эцэг эхийн боловсролын дундаж нь хэрэгсэл хувьсагч, харин 17-18 насны сурагч-багшийн дундаж харьцаа нь босго хувьсагчийн үүрэгтэй байсан. Энэхүү тайланд уг суурь логикийг хадгалж, өгөөжийн ялгаатай байдлыг нэг урьдчилан сонгосон босгоор хязгаарлахгүйгээр, нас, хүйс, гэрлэлтийн байдал, хотжилт, сургуулийн хүртээмж, сурагч-багшийн харьцаа, төрсөн үе, судалгааны давалгаа зэрэг шинжүүдтэй хамтатган үнэлэв.",
)
add_para(
    doc,
    f"Шинжилгээнд үндсэн ХХБР түүврийн {int(summary['N']):,} ажиглалтыг ашиглаж, {int(summary['num_trees']):,} модтой Instrumental Forest загвар байгуулсан. Тооцооны үр дүнгээр боловсролын нэг жилийн таамагласан local IV өгөөж дунджаар {pct(summary['mean_return_pct'])} хувь, медианаар {pct(summary['median_return_pct'])} хувь байна. Local IV өгөөжийн 10-90 хувийн муж {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байгаа нь боловсролын өгөөж нэг дундаж коэффициентоор бүрэн тайлбарлагдахгүй, бүлэг хоорондын ялгаатай байдлыг агуулж байгааг харуулна.",
)
add_para(
    doc,
    "Variable importance-ийн үр дүнгээс харахад өгөөжийн ялгааг ялгахад хувийн шинжүүд хамгийн өндөр ач холбогдолтой боловч сургуулийн хүртээмж болон сурагч-багшийн харьцаа мөн чухал мэдээлэл өгч байна. Иймээс Instrumental Forest-ийн үр дүн нь үндсэн ХХБР загварын босго үр дүнг орлох бус, харин боловсролын өгөөжийн ялгаа сургуулийн орчин болон хувь хүний шинжүүдтэй хавсарсан олон хэмжээст үзэгдэл болохыг нэмэлтээр харуулж байна.",
)
add_para(doc, "Түлхүүр үгс: боловсролын өгөөж; Instrumental Forest; хэрэгсэл хувьсагч; машин сургалт; сурагч-багшийн харьцаа.", first_line=False)
add_para(doc, "JEL ангилал: C14, C26, I26, J31.", first_line=False)

# Introduction
add_heading(doc, "ОРШИЛ", 1)
add_para(
    doc,
    "Боловсролын эдийн засгийн өгөөжийг үнэлэх судалгаанд дундаж коэффициент нь бодлогын ерөнхий чиглэлийг тодорхойлоход ашигтай боловч бүх бүлэгт адил үйлчилдэг гэж үзэхэд хэт ерөнхий дүгнэлт болох эрсдэлтэй. Хөдөлмөрийн зах зээл, сургуулийн чанар, гэр бүлийн суурь нөхцөл, хот хөдөөгийн ялгаа зэрэг хүчин зүйлс боловсролын нэг жилийн цалингийн өгөөжийг өөр өөрөөр илэрхийлж болно. Иймээс орчин үеийн хөдөлмөрийн эдийн засагт боловсролын өгөөжийн ялгаатай байдлыг буюу heterogeneous returns to schooling-ийг тусгайлан авч үзэх шаардлага нэмэгдэж байна.",
)
add_para(
    doc,
    "Үндсэн судалгаанд боловсролын өгөөжийг ЭХБК, ХШХБК болон ХХБР загваруудаар үнэлж, 17-18 насанд харгалзах сурагч-багшийн дундаж харьцаагаар тодорхойлогдох сургуулийн орчны босгоны ялгааг шалгасан. Тус үр дүн нь боловсролын өгөөж нэг шугаман дундаж тоогоор хязгаарлагдахгүй байж болохыг харуулсан боловч нэг босго сонголтод тулгуурласан арга тул нэмэлт эрэл хайгуулын шалгалт шаардана.",
)
add_para(
    doc,
    "Энэ нэмэлт тайлангийн зорилго нь Instrumental Forest аргыг ашиглан боловсролын өгөөжийн ялгаатай байдлыг өгөгдөлд суурилсан хэлбэрээр дүрслэх явдал юм. Ингэхдээ үндсэн шалтгаант тайлбарыг өөрчлөхгүй, харин ХХБР загвараар илэрсэн босгоны үр дүн олон шинжтэй хавсарсан heterogeneity-ийн нэг хэсэг мөн эсэхийг шалгах supplementary evidence гаргана.",
)
add_para(
    doc,
    "Судалгааны ажил таван үндсэн хэсгээс бүрдэнэ. Нэгдүгээр бүлэгт боловсролын өгөөж, heterogeneous treatment effect болон Instrumental Forest аргын судлагдсан байдлыг товч танилцуулна. Хоёрдугаар бүлэгт загварын арга зүй, хувьсагчдын үүргийг тайлбарлана. Гуравдугаар бүлэгт өгөгдлийн бүтэц, шинжилгээний түүврийг танилцуулна. Дөрөвдүгээр бүлэгт local IV өгөөжийн тархалт, STR-тэй уялдах дүр зураг, variable importance-ийн үр дүнг тайлбарлана. Тавдугаар бүлэгт уг нэмэлт шинжилгээг үндсэн тайланд хэрхэн болгоомжтой ашиглах зөвлөмжийг нэгтгэнэ.",
)

# Literature
add_heading(doc, "I БҮЛЭГ. СУДЛАГДСАН БАЙДАЛ", 1)
add_heading(doc, "1.1. Боловсролын өгөөж ба ялгаатай нөлөө", 2)
add_para(
    doc,
    "Боловсролын өгөөжийн судалгааны суурь нь Минсерийн цалингийн тэгшитгэл бөгөөд боловсролын нэмэлт нэг жил хөдөлмөрийн орлоготой хэрхэн холбогдохыг үнэлдэг (Mincer, 1974). Гэвч боловсролын сонголт санамсаргүй бус байдаг тул ЭХБК үнэлгээ нь ажиглагдахгүй чадвар, гэр бүлийн суурь нөхцөл, боловсролын хэмжилтийн алдаанаас шалтгаалан хазайлттай байж болно. Card (1999) боловсролын өгөөжийг үнэлэхэд хэрэгсэл хувьсагчийн арга чухал боловч үр дүнг identification-ийн таамаглалтай хамт болгоомжтой тайлбарлах шаардлагатайг онцолсон.",
)
add_para(
    doc,
    "Боловсролын өгөөж зөвхөн дундаж түвшинд бус, тухайн хүний сурсан орчин, хөдөлмөрийн зах зээлийн боломж, хотжилт, гэр бүлийн боловсролын түвшин зэргээс хамааран ялгаатай илэрч болно. Энэ санаа нь олон улсын боловсролын эдийн засгийн судалгаанд heterogeneous treatment effect гэсэн ерөнхий ойлголтоор хөгжсөн бөгөөд бодлогын дүгнэлт хийхэд дундаж өгөөжөөс гадна хэнд илүү өндөр, хэнд бага өгөөж илэрч буйг харах шаардлагатай болгодог.",
)
add_heading(doc, "1.2. Machine learning ба causal forest чиглэлийн хэрэглээ", 2)
add_para(
    doc,
    "Athey, Tibshirani, and Wager (2019) нар causal forest аргыг боловсруулж, treatment effect-ийн ялгаатай байдлыг decision tree болон random forest-ийн санаанд тулгуурлан үнэлэх боломжийг харуулсан. Chernozhukov et al. (2018)-ийн double/debiased machine learning чиглэл нь nuisance function-уудыг machine learning-ээр тооцоолохын зэрэгцээ шалтгаант параметрийг хамгаалалттай үнэлэх арга зүйн суурийг өргөжүүлсэн. Instrumental Forest нь энэ урсгалын хүрээнд эндоген тайлбарлагчтай нөхцөлд локал IV өгөөжийг өгөгдлийн шинжүүдээр ялган харах боломж олгодог.",
)
add_para(
    doc,
    "Энэхүү тайланд Instrumental Forest-ийг үндсэн causal proof болгон бус, ХХБР загварын нэмэлт шалгалт болгон ашиглаж байна. Ингэснээр нэг босгоны сонголтоос үүдэх хэт хамаарлыг бууруулж, боловсролын өгөөжийн ялгаа хувь хүний болон сургуулийн орчны олон шинжтэй хамт илэрч байгаа эсэхийг дүрслэх боломж бүрдэнэ.",
)

# Method
add_heading(doc, "II БҮЛЭГ. АРГА ЗҮЙ", 1)
add_heading(doc, "2.1. Загварын үндсэн логик", 2)
add_para(
    doc,
    "Instrumental Forest-ийн үндсэн зорилго нь боловсролын нэг жилийн өгөөжийг нэг дундаж коэффициентээр хязгаарлахгүй, харин ажиглагдсан шинжүүдийн нөхцөлд ялгаатай байдлаар таамаглах явдал юм. Судалгаанд дараах тэмдэглэгээг ашиглав.",
)
add_para(doc, "Y_i = ln(wage_i), W_i = educ_years_i, Z_i = parent_educ_mean_i, X_i = observed characteristics_i. [1]", WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
add_para(
    doc,
    "Энд Y_i нь хувь хүний бодит цагийн цалингийн логарифм, W_i нь боловсролын жил, Z_i нь эцэг эхийн боловсролын дундаж буюу хэрэгсэл хувьсагч, X_i нь нас, хүйс, гэрлэлтийн байдал, хотжилт, сурагч-багшийн харьцаа, сургуулийн хүртээмж, төрсөн аймаг, төрсөн үе, судалгааны давалгаа зэрэг heterogeneity-г тайлбарлах шинжүүд юм.",
)
add_para(
    doc,
    "Загварын тайлбарлах параметрийг дараах байдлаар ойлгоно.",
)
add_para(doc, "tau(X_i) = local IV return to an additional year of schooling conditional on X_i. [2]", WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
add_para(
    doc,
    "Энэ tau(X_i)-г тухайн хүний яг хувийн causal effect гэж бус, ижил төстэй шинжтэй бүлгийн боловсролын нөхцөлт IV өгөөж гэж тайлбарлана. Иймээс тайлангийн бүх үр дүн exploratory буюу эрэл хайгуулын шинжтэй байна.",
)
add_heading(doc, "2.2. Үндсэн ХХБР загвартай холбогдох байдал", 2)
add_para(
    doc,
    "Үндсэн тайлангийн ХХБР загвар 17-18 насны сурагч-багшийн дундаж харьцааны 19.53 орчим босгоор боловсролын өгөөж ялгаатай эсэхийг шалгасан. Instrumental Forest нь уг босгыг урьдчилан гол тайлбарлагч болгон тулгахгүй, харин сурагч-багшийн харьцааг X шинжийн нэг хэсэг болгон оруулж, бусад шинжүүдтэй хамт өгөөжийн ялгааг дүрслэнэ. Энэ нь ХХБР-ийн үр дүнг батлах эсвэл үгүйсгэх ганц тест биш, харин heterogeneity-ийн илүү өргөн зураглал юм.",
)
add_caption(doc, "Хүснэгт 1. Instrumental Forest шинжилгээний загварын бүтэц")
add_table(
    doc,
    ["Элемент", "Хувьсагч", "Тайлбар"],
    [
        ["Хамааруулах хувьсагч", "lwage", "Бодит цагийн цалингийн логарифм"],
        ["Эндоген тайлбарлагч", "educ_years", "Боловсролын нийт жил"],
        ["Хэрэгсэл хувьсагч", "parent_educ_mean", "Эцэг эхийн боловсролын дундаж"],
        ["Гол орчны шинж", "student_teacher_ratio_avg_17_18", "17-18 насны сурагч-багшийн дундаж харьцаа"],
        ["Нэмэлт шинжүүд", "age, female, married, urban, q_school_access, FE dummies", "Өгөөжийн ялгааг сургах X матриц"],
        ["Загвар", "grf::instrumental_forest", f"{int(summary['num_trees']):,} модтой локал IV forest"],
    ],
    [4.0, 5.2, 6.3],
)
add_source(doc)

# Data
add_heading(doc, "III БҮЛЭГ. ӨГӨГДӨЛ БА ХУВЬСАГЧИД", 1)
add_para(
    doc,
    f"Шинжилгээнд үндсэн ХХБР загварын эцсийн түүвэр болох {int(summary['N']):,} ажиглалтыг ашиглав. Энэ түүвэр нь боловсрол, цалин, эцэг эхийн боловсрол, төрсөн аймаг, 17-18 насанд харгалзах сургуулийн орчны үзүүлэлт бүрэн ажиглагдсан ажиллагчдаас бүрдэнэ. Иймээс Instrumental Forest-ийн түүвэр үндсэн ХХБР үр дүнтэй шууд харьцуулах боломжтой.",
)
add_para(
    doc,
    "X матрицад нийт 48 шинж орсон. Үүнд хувь хүний демограф шинжүүд, сургуулийн орчны үзүүлэлтүүд, төрсөн аймаг, төрсөн үе, судалгааны давалгааны дамми хувьсагчид багтсан. Харин parent_educ_mean нь зөвхөн хэрэгсэл хувьсагчийн үүрэгтэй тул X матрицад тайлбарлагч шинж болгон оруулаагүй. Энэ нь хэрэгсэл хувьсагчийн үүргийг threshold эсвэл heterogeneity feature-тэй хольж тайлбарлах эрсдэлийг бууруулна.",
)

# Results
add_heading(doc, "IV БҮЛЭГ. ШИНЖИЛГЭЭНИЙ ҮР ДҮН", 1)
add_heading(doc, "4.1. Local IV өгөөжийн тархалт", 2)
add_caption(doc, "Хүснэгт 2. Local IV өгөөжийн үндсэн статистик")
add_table(
    doc,
    ["Үзүүлэлт", "Утга"],
    [
        ["Ажиглалтын тоо", f"{int(summary['N']):,}"],
        ["Модны тоо", f"{int(summary['num_trees']):,}"],
        ["X шинжийн тоо", f"{int(summary['X_features'])}"],
        ["Дундаж local IV өгөөж", f"{pct(summary['mean_return_pct'])}%"],
        ["Медиан local IV өгөөж", f"{pct(summary['median_return_pct'])}%"],
        ["10 хувийн перцентиль", f"{pct(summary['p10_return_pct'])}%"],
        ["90 хувийн перцентиль", f"{pct(summary['p90_return_pct'])}%"],
        ["Харьцуулах ХШХБК өгөөж", f"{pct(100*(2.718281828459045**summary['IV_2SLS_beta_log_reference']-1))}%"],
    ],
    [8.0, 5.0],
)
add_source(doc)
add_para(
    doc,
    f"Хүснэгт 2-оос харахад Instrumental Forest-ийн таамагласан local IV өгөөж дунджаар {pct(summary['mean_return_pct'])} хувь, медианаар {pct(summary['median_return_pct'])} хувь байна. Энэ нь үндсэн ХШХБК үнэлгээнээс гарсан дундаж өгөөжтэй ойролцоо тул model нийт түвшинд судалгааны суурь үр дүнгээс тасарч, хэт өөр дүр зураг гаргаагүй байна. Харин 10-90 хувийн муж {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байгаа нь боловсролын өгөөж бүлгүүдээр нэлээд ялгаатай тархаж болохыг харуулж байна.",
)
add_caption(doc, "Зураг 1. Instrumental Forest өгөөжийн тархалт")
add_picture(doc, FIG_DIR / "instrumental_forest_tau_distribution.png")
add_source(doc)
add_para(
    doc,
    "Зураг 1-д local IV өгөөжийн тархалтыг харуулав. Тасархай босоо шугам нь ХШХБК-ийн дундаж өгөөжийг илэрхийлнэ. Тархалтын төв хэсэг 10-12 хувийн орчимд байрлаж байгаа боловч баруун талдаа өндөр өгөөжтэй бүлгүүд ажиглагдаж байна. Иймээс боловсролын өгөөжийг зөвхөн нэг дундаж коэффициентээр тайлбарлах нь мэдээлэл алдах боломжтой.",
)

add_heading(doc, "4.2. Өгөөжийн квартил бүлгүүд", 2)
add_caption(doc, "Хүснэгт 3. Local IV өгөөжийн квартил бүлгүүдийн харьцуулалт")
add_table(
    doc,
    ["Бүлэг", "N", "Дундаж өгөөж", "STR медиан", "Хотын хувь", "Эмэгтэй хувь"],
    [
        [
            r["tau_group_label"],
            f"{int(float(r['N'])):,}",
            f"{pct(float(r['mean_return_pct']))}%",
            num(float(r["median_student_teacher_ratio_17_18"])),
            f"{pct(100*float(r['urban_share']))}%",
            f"{pct(100*float(r['female_share']))}%",
        ]
        for r in quartiles
    ],
    [3.3, 2.0, 2.8, 2.5, 2.4, 2.4],
)
add_source(doc)
add_para(
    doc,
    "Квартилын харьцуулалтаас өгөөжийн доод 25 хувийн бүлэгт дундаж өгөөж 6.95 хувь, дээд 25 хувийн бүлэгт 16.93 хувь байна. Энэ ялгаа нь Instrumental Forest-ийн гол нэмэлт мэдээлэл юм. Өөрөөр хэлбэл, боловсролын нэг жилийн өгөөж нь судалгааны бүх хүний хувьд адил биш бөгөөд хувь хүний шинж, сургуулийн орчин, судалгааны давалгааны ялгаатай хамт илэрч байна.",
)

add_heading(doc, "4.3. STR босготой холбож унших нь", 2)
add_caption(doc, "Хүснэгт 4. STR босгоны регимээр харьцуулсан Instrumental Forest үр дүн")
add_table(
    doc,
    ["Регим", "N", "Дундаж өгөөж", "STR медиан", "Хотын хувь", "Дундаж нас"],
    [
        [
            r["threshold_regime"],
            f"{int(float(r['N'])):,}",
            f"{pct(float(r['mean_return_pct']))}%",
            num(float(r["median_student_teacher_ratio_17_18"])),
            f"{pct(100*float(r['urban_share']))}%",
            num(float(r["mean_age"])),
        ]
        for r in regimes
    ],
    [4.0, 2.0, 2.8, 2.7, 2.5, 2.3],
)
add_source(doc)
add_para(
    doc,
    "STR-ийн 19.53 босгоор хоёр регимд хувааж дундажлахад Instrumental Forest-ийн local IV өгөөж доод регимд 11.99 хувь, дээд регимд 11.04 хувь байна. Энэ үр дүн нь ХХБР загварын дээд STR регимийн өндөр өгөөжийг шууд давтаагүй. Гэхдээ энэ нь загвар буруу гэсэн үг биш, харин өгөөжийн ялгаа зөвхөн нэг босгоор тайлбарлагдахгүй, бусад шинжүүдтэй хамт нөхцөлдөн илэрч байгааг харуулж байна.",
)
add_caption(doc, "Зураг 2. Local IV өгөөж ба сурагч-багшийн харьцааны хамаарал")
add_picture(doc, FIG_DIR / "instrumental_forest_tau_by_student_teacher_ratio.png")
add_source(doc)
add_para(
    doc,
    "Зураг 2-т local IV өгөөжийг 17-18 насны сурагч-багшийн дундаж харьцаатай харьцуулан дүрслэв. Босго орчимд хоёр регимийн ажиглалтууд нягт байрлаж байгаа бөгөөд өгөөжийн муруй шулуун өсөлт эсвэл бууралтын нэг хэв маягтай биш байна. Иймээс STR нь өгөөжийн ялгааг тайлбарлахад оролцож байгаа боловч ганцаараа бүх ялгааг тодорхойлох үзүүлэлт биш гэж тайлбарлах нь зүйтэй.",
)

add_heading(doc, "4.4. Variable importance-ийн үр дүн", 2)
add_caption(doc, "Хүснэгт 5. Variable importance-ийн нэгтгэл")
add_table(
    doc,
    ["Хувьсагчийн бүлэг", "Харьцангуй ач холбогдол"],
    [[r["family"], f"{pct(100*float(r['share']))}%"] for r in importance],
    [8.0, 5.0],
)
add_source(doc)
add_caption(doc, "Зураг 3. Өгөөжийн ялгааг тайлбарлах хувьсагчдын ач холбогдол")
add_picture(doc, FIG_DIR / "instrumental_forest_variable_importance.png", width_inches=6.1)
add_source(doc)
add_para(
    doc,
    "Variable importance-ийн үр дүнгээр өгөөжийн ялгааг ялгахад хувийн шинжүүд 48.02 хувь, сургуулийн хүртээмж 19.23 хувь, сурагч-багшийн харьцаа 15.23 хувь, судалгааны давалгаа 10.22 хувь, төрсөн үе 7.29 хувийн харьцангуй ач холбогдолтой байна. Энэ нь боловсролын өгөөжийн ялгаа зөвхөн STR босгоор биш, хувийн болон институцийн олон шинжээр хамт тодорхойлогдож байгааг харуулна.",
)

# Conclusion
add_heading(doc, "V БҮЛЭГ. ДҮГНЭЛТ БА ТАЙЛАНД АШИГЛАХ ЗӨВЛӨМЖ", 1)
add_para(
    doc,
    f"Энэхүү нэмэлт тайлангаар Instrumental Forest аргыг ашиглан боловсролын өгөөжийн ялгаатай байдлыг эрэл хайгуулын байдлаар үнэлэв. Үр дүнгээр local IV өгөөжийн дундаж {pct(summary['mean_return_pct'])} хувь, медиан {pct(summary['median_return_pct'])} хувь гарсан нь үндсэн ХШХБК үнэлгээтэй ойролцоо байна. Иймээс model нийт өгөөжийн түвшинд үндсэн судалгаатай нийцэж байна.",
)
add_para(
    doc,
    f"Гэхдээ local IV өгөөжийн 10-90 хувийн муж {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байгаа нь боловсролын өгөөж нэг дундаж коэффициентээр бүрэн тайлбарлагдахгүйг харуулна. Өгөөжийн дээд 25 хувийн бүлгийн дундаж өгөөж 16.93 хувь, доод 25 хувийн бүлгийнх 6.95 хувь байгаа нь subgroup heterogeneity бодитой байж болохыг илтгэнэ.",
)
add_para(
    doc,
    "Энэ шинжилгээний хамгийн зөв хэрэглээ нь үндсэн ХХБР үр дүнг орлуулах бус, түүнийг баяжуулах явдал юм. ХХБР загвар 17-18 насны сурагч-багшийн харьцааны тодорхой босгоор өгөөж ялгаатай эсэхийг шалгадаг бол Instrumental Forest нь өгөөжийн ялгаа сургуулийн хүртээмж, STR, хувь хүний шинж, судалгааны давалгаа, төрсөн үе зэрэг олон хүчин зүйлтэй хамт илэрч байгааг харуулна.",
)
add_para(
    doc,
    "Иймээс үндсэн тайланд дараах болгоомжтой өгүүлбэрийг ашиглахыг зөвлөж байна: Instrumental Forest-ийн нэмэлт шинжилгээ нь боловсролын өгөөжийн дундаж таамаглал ХШХБК үнэлгээтэй ойролцоо боловч өгөөж хувь хүмүүсийн шинж болон сургуулийн орчны үзүүлэлтүүдээр ялгаатай тархаж байгааг харууллаа. Энэхүү үр дүн нь ХХБР загварын босго үр дүнг шалтгаант баталгаа болгон орлохгүй, харин боловсролын өгөөжийн heterogeneity олон хэмжээст шинжтэй болохыг дэмжсэн exploratory evidence юм.",
)
add_para(
    doc,
    "Харин дараах дүгнэлтийг хийхээс зайлсхийх шаардлагатай: Instrumental Forest нь STR-ийн дээд регимд өгөөж статистикийн хувьд заавал өндөр гэдгийг баталсан гэж хэлж болохгүй. Учир нь forest-ийн регимээр дундажласан үр дүн ХХБР-ийн чиглэлийг яг давтаагүй бөгөөд энэ арга нь статистикийн ач холбогдлын уламжлалт нэг p-утгатай тест бус, heterogeneity-г зураглах нэмэлт machine learning арга юм.",
)

# References
add_heading(doc, "НОМ ЗҮЙ", 1)
refs = [
    "Athey, S., Tibshirani, J., & Wager, S. (2019). Generalized random forests. Annals of Statistics, 47(2), 1148-1178.",
    "Card, D. (1999). The causal effect of education on earnings. In O. Ashenfelter & D. Card (Eds.), Handbook of Labor Economics (Vol. 3A, pp. 1801-1863). Elsevier.",
    "Caner, M., & Hansen, B. E. (2004). Instrumental variable estimation of a threshold model. Econometric Theory, 20(5), 813-843.",
    "Chernozhukov, V., Chetverikov, D., Demirer, M., Duflo, E., Hansen, C., Newey, W., & Robins, J. (2018). Double/debiased machine learning for treatment and structural parameters. Econometrics Journal, 21(1), C1-C68.",
    "Mincer, J. (1974). Schooling, experience, and earnings. National Bureau of Economic Research.",
    "Tibshirani, J., Athey, S., Friedberg, R., Hadad, V., Hirshberg, D., Miner, L., Sverdrup, E., Wager, S., Wright, M., & grf contributors. (2024). grf: Generalized Random Forests. R package.",
    "Үндэсний статистикийн хороо. (2024). Өрхийн нийгэм эдийн засгийн судалгаа 2024. ҮСХ. https://data.1212.mn",
]
for i, ref in enumerate(refs, 1):
    p = add_para(doc, f"{i}. {ref}", first_line=False)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)

# Appendix
add_heading(doc, "ХАВСРАЛТ", 1)
add_heading(doc, "Хавсралт А. Давтан ажиллуулах код", 2)
add_para(
    doc,
    "Instrumental Forest шинжилгээг дараах R script-ээр бүрэн давтан ажиллуулах боломжтой: R/31_instrumental_forest.R. Script нь ivtr_ready_parent_educ_mean_student_teacher_avg_17_18.rds түүврийг уншиж, grf::instrumental_forest функцээр local IV өгөөжийг тооцоод output/tables болон output/figures хавтаст хүснэгт, зургуудыг хадгална.",
)
add_para(doc, "PowerShell жишээ команд:", first_line=False)
add_para(doc, "$env:GRF_TREES='2000'; $env:GRF_THREADS='3'; Rscript R/31_instrumental_forest.R", first_line=False)
add_heading(doc, "Хавсралт Б. Үр дүнгийн файлууд", 2)
for item in [
    "output/tables/T16_instrumental_forest_summary.csv",
    "output/tables/T16_instrumental_forest_tau_quartiles.csv",
    "output/tables/T16_instrumental_forest_threshold_regimes.csv",
    "output/tables/T16_instrumental_forest_variable_importance.csv",
    "output/figures/instrumental_forest_tau_distribution.png",
    "output/figures/instrumental_forest_tau_by_student_teacher_ratio.png",
    "output/figures/instrumental_forest_variable_importance.png",
]:
    add_para(doc, item, first_line=False)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
