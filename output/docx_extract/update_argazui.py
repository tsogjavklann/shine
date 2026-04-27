from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOCX_PATH = Path("bagiin_ner_v29.docx")
BACKUP_DIR = Path("output/docx_backups")
FALLBACK_PATH = Path("output/bagiin_ner_v29_argazui_updated.docx")

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


FORMULAS = {
    "2": r"\ln(w_i)=\alpha+\beta\,educ_i+X_i^{\prime}\delta+\lambda_{a(i)}+\kappa_{c(i)}+\tau_{t(i)}+\varepsilon_i",
    "3": r"\hat{\theta}_{OLS}=\arg\min_{\theta}\sum_{i=1}^{N}h_i\left[\ln(w_i)-\left(\alpha+\beta\,educ_i+X_i^{\prime}\delta+\lambda_{a(i)}+\kappa_{c(i)}+\tau_{t(i)}\right)\right]^2",
    "4": r"r(\hat{\beta})=100\left[\exp(\hat{\beta})-1\right]",
    "5": r"educ_i=\pi_0+\pi_1 z_i+X_i^{\prime}\pi_2+\lambda_{a(i)}+\kappa_{c(i)}+\tau_{t(i)}+v_i",
    "6": r"\ln(w_i)=\alpha+\beta\,\widehat{educ}_i+X_i^{\prime}\delta+\lambda_{a(i)}+\kappa_{c(i)}+\tau_{t(i)}+u_i",
    "7": r"\ln(w_i)=\mathbb{1}(q_i\leq\gamma)(\alpha_1+\beta_1\,educ_i+X_i^{\prime}\delta_1)+\mathbb{1}(q_i>\gamma)(\alpha_2+\beta_2\,educ_i+X_i^{\prime}\delta_2)+\lambda_{a(i)}+\kappa_{c(i)}+\tau_{t(i)}+u_i",
    "8": r"\hat{\gamma}=\arg\min_{\gamma\in\Gamma}SSR(\gamma),\quad \Gamma=[Q_{0.15}(q),Q_{0.85}(q)],\quad \Delta_{\beta}=\hat{\beta}_{high}-\hat{\beta}_{low}",
}


METHODOLOGY_ITEMS = [
    ("text", "Энэхүү судалгааны арга зүй нь боловсролын цалингийн өгөөжийг үнэлэхдээ гурван шатны логик дараалал баримтална. Нэгдүгээрт, Минсерийн уламжлалт цалингийн тэгшитгэлийг ЭХБК аргаар үнэлж суурь холбоог тогтооно. Хоёрдугаарт, боловсролын жил санамсаргүй сонгогддоггүй гэсэн эндоген асуудлыг эцэг, эхийн боловсролын дундаж хэрэгсэл хувьсагчаар засаж ХШХБК үнэлгээ хийнэ. Гуравдугаарт, боловсролын өгөөж сургуулийн орчны ялгаагаар өөр өөр байж болох эсэхийг Caner, Hansen (2004)-ийн санаанд тулгуурласан хэрэгсэл хувьсагчтай босго утгат регрессээр шалгана. Үндсэн шинжилгээний нэгж нь ӨНЭЗС-ийн 2020, 2021, 2022, 2024 оны давалгаанд хамрагдсан 25–60 насны цалин хөлстэй ажиллагч бөгөөд загваруудад өрхийн түүврийн жин, төрсөн аймаг, төрсөн үеийн бүлэг, судалгааны давалгааны тогтмол нөлөөг харгалзан үзнэ."),
    ("text", "2.1. Судалгааны загвар ба хувьсагчдын тогтолцоо"),
    ("text", "Судалгааны хамаарах хувьсагч нь сарын хөдөлмөрийн орлогын логарифмчилсан утга байна. Үндсэн тайлбарлагч нь боловсролын жил, хяналтын хувьсагчид нь нас, насны квадрат, хүйс, гэрлэлтийн байдал, хотын суурьшлын үзүүлэлтээс бүрдэнэ. Тогтмол нөлөөгөөр төрсөн аймаг, төрсөн үеийн бүлэг, судалгааны давалгааг оруулснаар орон нутгийн суурь ялгаа, үеийн ялгаа, тухайн жилийн макро орчны нөлөөг хянах боломжтой болно. Суурь бүтцийн тэгшитгэлийг дараах байдлаар бичнэ:"),
    ("equation", "2"),
    ("text", "Энд зүүн талын хувьсагч нь логаритмчилсан цалин, боловсролын хувьсагч нь боловсролын жил, хяналтын вектор нь нас, насны квадрат, хүйс, гэрлэлтийн байдал, хотын суурьшлыг илэрхийлнэ. Тэгшитгэлд төрсөн аймаг, төрсөн үеийн бүлэг, судалгааны давалгааны тогтмол нөлөөг тусгасан. Боловсролын коэффициент нь нэмэлт нэг жилийн цалингийн хагас уян хатан чанарыг илэрхийлэх бөгөөд бүх стандарт алдааг төрсөн аймгийн түвшинд кластерлан тооцно."),
    ("text", "2.2. ЭХБК суурь үнэлгээ"),
    ("text", "ЭХБК үнэлгээ нь боловсрол ба цалингийн хоорондын суурь статистик холбоог харуулах зорилготой. Гэхдээ энэ үнэлгээг шууд шалтгаант нөлөө гэж тайлбарлахгүй. Учир нь боловсролын сонголт хувь хүний ажиглагдашгүй ур чадвар, өрхийн боловсролын орчин, хөдөлмөрийн зах зээлийн мэдээлэл авах боломж зэрэг хүчин зүйлтэй холбоотой байж болно. Иймээс ЭХБК нь дараагийн ХШХБК болон ХХБР үнэлгээтэй харьцуулах суурь үзүүлэлт болно."),
    ("equation", "3"),
    ("text", "Энд h_i нь ӨНЭЗС-ийн өрхийн түүврийн жин, θ нь үнэлэгдэх бүх параметрийн вектор юм. Жинлэсэн ЭХБК үнэлгээ нь түүврийн төлөөлөх чадварыг хадгалж, кластерлагдсан стандарт алдаа нь нэг төрсөн аймагт харьяалагдах ажиглалтуудын хоорондын хамаарлыг зөвшөөрнө."),
    ("equation", "4"),
    ("text", "Логаритмчилсан цалингийн загварт боловсролын коэффициентийг хувийн өгөөж болгон тайлбарлахдаа экспоненциал хувиргалтыг хэрэглэнэ. Иймд үр дүнгийн бүлэгт боловсролын нэг жилийн өгөөжийг зөвхөн коэффициентээр бус, хувиар илэрхийлсэн өгөөжөөр тайлбарлана."),
    ("text", "2.3. Хэрэгсэл хувьсагчтай үнэлгээ"),
    ("text", "Боловсролын жил эндоген байж болох тул үндсэн шалтгаант үнэлгээнд ХШХБК аргыг хэрэглэнэ. Энэ судалгаанд эцэг, эхийн боловсролын дундажийг боловсролын жилийн хэрэгсэл хувьсагчаар сонгосон. Уг хувьсагч хүүхдийн боловсрол үргэлжлүүлэх боломж, суралцах орчин, өрхийн боловсролын үнэлэмжтэй холбоотой тул боловсролын жилтэй хүчтэй хамаарах онолын үндэслэлтэй. Харин цалинд гэр бүлийн нийгмийн сүлжээ, хүмүүжил зэрэг сувгаар шууд нөлөөлөх боломжийг бүрэн үгүйсгэх боломжгүй тул үр дүнг болгоомжтой, тодорхой таамаглалын хүрээнд тайлбарлана."),
    ("text", "ХШХБК үнэлгээ хоёр шаттай. Эхний шатанд боловсролын жилийг хэрэгсэл хувьсагч болон бусад хяналтын хувьсагчаар тайлбарлаж, боловсролын эндоген бус хэсгийг ялгана. Энэ шат нь хэрэгсэл хувьсагчийн хамаарал буюу relevance нөхцөлийг шалгах үндсэн суурь болно:"),
    ("text", "Эхний шатны тэгшитгэлийг дараах байдлаар тодорхойлно:"),
    ("equation", "5"),
    ("text", "Энд хэрэгсэл хувьсагч нь эцэг, эхийн боловсролын дундаж юм. Хэрэгсэл хувьсагч сул эсэхийг эхний шатны F статистикаар үнэлэх бөгөөд F утга 10-аас их байх ерөнхий шаардлагыг хангах эсэхийг шалгана. Энэ судалгаанд нэг эндоген хувьсагч, нэг хэрэгсэл хувьсагч ашиглаж байгаа тул загвар яг тодорхойлогдсон хэлбэртэй бөгөөд хэт тодорхойлогдсон байдлын J шалгуурыг хэрэглэхгүй."),
    ("equation", "6"),
    ("text", "Хоёр дахь шатанд эхний шатнаас гарсан боловсролын тохирсон утгыг үндсэн цалингийн тэгшитгэлд ашиглана. Ингэснээр β коэффициент нь эцэг, эхийн боловсролоор өдөөгдсөн боловсролын жилийн өөрчлөлтийн хүрээнд тодорхойлогдох өгөөжийг илэрхийлнэ."),
    ("text", "2.4. Хэрэгсэл хувьсагчтай босго утгат регресс (ХХБР)"),
    ("text", "Судалгааны үндсэн арга нь Caner, Hansen (2004)-ийн хэрэгсэл хувьсагчтай босго утгат регресс юм. Энэ арга нь хоёр асуудлыг нэгэн зэрэг авч үздэг: боловсролын жил эндоген байж болох асуудал болон боловсролын өгөөж нэг ижил налуутай бус, тодорхой босго хувьсагчаар ялгарах регимүүдийн хооронд өөр байж болох асуудал. Энэхүү судалгаанд босго хувьсагч нь боловсролын жил биш, харин тухайн хүний 17–18 насанд харгалзах сурагч-багшийн дундаж харьцаа юм."),
    ("equation", "7"),
    ("text", "Энд босго хувьсагч нь 17–18 насны сурагч-багшийн дундаж харьцаа, харин босго утга нь өгөгдлөөс үнэлэгдэх параметр юм. Босгоноос бага эсвэл тэнцүү харьцаатай үед багшийн ачаалал харьцангуй бага сургуулийн орчны регим, босгоноос дээш харьцаатай үед багшийн ачаалал харьцангуй өндөр сургуулийн орчны регим гэж ялгана. Боловсролын жил нь хоёр регимд хоёуланд нь эндоген гэж үзэгдэх тул эцэг, эхийн боловсролын дундаж хэрэгсэл хувьсагчийн үүрэг хэвээр хадгалагдана."),
    ("text", "Босгыг урьдчилан оноож сонгохгүй, харин боломжит утгуудын сүлжээн хайлтаар тодорхойлно. Хэт бага ажиглалттай регим үүсэхээс сэргийлж босгоны хайлтыг босго хувьсагчийн 15–85 хувийн мужид хязгаарлана. Боломжит босго бүр дээр хоёр регимийн ХШХБК үнэлгээг хийж, нэгтгэн агшаасан алдааны квадратын нийлбэр хамгийн бага байх утгыг оновчтой босго болгон сонгоно."),
    ("equation", "8"),
    ("text", "Регимүүдийн ялгааг хоёр регимийн боловсролын өгөөжийн зөрүүгээр хэмжинэ. Босго болон регимийн ялгааны статистик ач холбогдлыг төрсөн аймгийн түвшинд кластерласан бүүтстрап дахин түүвэрлэлтийн аргаар үнэлнэ. Үндсэн шинжилгээнд 399 удаагийн бүүтстрап давталт ашиглаж, итгэлийн интервал болон p-утгыг эмпирик тархалтаас тооцно."),
    ("text", "Caner, Hansen (2004)-ийн алгоритмын бэлэн хэрэгжүүлэлт R болон Python-ийн түгээмэл багцуудад бүрэн хэлбэрээр байхгүй тул энэ судалгаанд тогтмол нөлөөг үлдэгдэлжүүлсэн, жинлэсэн, төрсөн аймгаар кластерласан гарын авлагын хэрэгжүүлэлт ашиглав. Иймд аргачлалыг эх зохиогчдын санаанд нийцүүлсэн ХХБР хэрэгжүүлэлт гэж тодорхойлж, босго хувьсагчийг цалингийн шууд шалтгаан бус, өгөөжийн ялгаатай регимийг ялгах шалгуур гэж тайлбарлана."),
    ("text", "2.5. Үр дүнгийн бат бөх байдлын шинжилгээний бүтэц"),
    ("text", "Үр дүнгийн бат бөх байдлыг хэд хэдэн түвшинд шалгана. Нэгдүгээрт, ЭХБК, ХШХБК, ХХБР үнэлгээнүүдийн боловсролын өгөөжийн хэмжээ, чиглэл, ач холбогдлыг харьцуулна. Хоёрдугаарт, эхний шатны F статистикаар эцэг, эхийн боловсролын дундаж хэрэгсэл хувьсагчийн хүчийг шалгана. Гуравдугаарт, босгоны сүлжээн хайлт болон бүүтстрап тархалтаар γ утгын тогтвортой байдлыг үнэлнэ. Дөрөвдүгээрт, регимүүдийн өгөөжийн ялгааны итгэлийн интервал тэгийг агуулж байгаа эсэхийг шалгана. Эдгээр шалгалт нь боловсролын өгөөжийн ялгаа зөвхөн ЭХБК-ийн хазайлт, сул хэрэгсэл хувьсагч, эсвэл санамсаргүй босго сонголтоос үүсээгүй гэдгийг харуулахад чиглэнэ."),
]


def normalized(text: str) -> str:
    return " ".join(text.split()).upper()


def generate_math_elements() -> dict[str, object]:
    md_lines = []
    for number, formula in FORMULAS.items():
        md_lines.append(f"$$\n{formula}\n$$")
    markdown = "\n\n".join(md_lines)

    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / "math.docx"
        result = subprocess.run(
            ["pandoc", "-f", "markdown+tex_math_dollars", "-o", str(out_path)],
            input=markdown,
            text=True,
            encoding="utf-8",
            capture_output=True,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(result.stderr)

        math_doc = Document(out_path)
        math_paragraphs = [p for p in math_doc.paragraphs if "oMath" in p._p.xml]
        if len(math_paragraphs) != len(FORMULAS):
            raise RuntimeError(f"Expected {len(FORMULAS)} formulas, found {len(math_paragraphs)}")

        elements = {}
        for number, paragraph in zip(FORMULAS.keys(), math_paragraphs):
            math_el = next(el for el in paragraph._p.iter() if el.tag == f"{{{M_NS}}}oMath")
            elements[number] = deepcopy(math_el)
        return elements


def ensure_equation_tabs(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tabs = p_pr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        p_pr.append(tabs)
    if not tabs.findall(qn("w:tab")):
        center = OxmlElement("w:tab")
        center.set(qn("w:val"), "center")
        center.set(qn("w:pos"), "4535")
        right = OxmlElement("w:tab")
        right.set(qn("w:val"), "right")
        right.set(qn("w:pos"), "9071")
        tabs.append(center)
        tabs.append(right)


def tab_run() -> object:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    return run


def number_run(number: str) -> object:
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:tab"))
    text = OxmlElement("w:t")
    text.text = f"({number})"
    run.append(text)
    return run


def replace_with_equation(paragraph, math_el, number: str) -> None:
    ensure_equation_tabs(paragraph)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph._p.append(tab_run())
    paragraph._p.append(deepcopy(math_el))
    paragraph._p.append(number_run(number))


def main() -> None:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = BACKUP_DIR / f"bagiin_ner_v29_before_argazui_{stamp}.docx"
    shutil.copy2(DOCX_PATH, backup)

    math_elements = generate_math_elements()
    doc = Document(DOCX_PATH)

    start = next(i for i, p in enumerate(doc.paragraphs) if normalized(p.text).startswith("II БҮЛЭГ"))
    end = next(
        i for i in range(start + 1, len(doc.paragraphs))
        if normalized(doc.paragraphs[i].text).startswith("III БҮЛЭГ")
    )
    slots = [i for i in range(start + 1, end) if doc.paragraphs[i].text.strip()]
    if len(slots) != len(METHODOLOGY_ITEMS):
        raise RuntimeError(f"Expected {len(METHODOLOGY_ITEMS)} methodology slots, found {len(slots)}")

    for idx, item in zip(slots, METHODOLOGY_ITEMS):
        paragraph = doc.paragraphs[idx]
        if item[0] == "equation":
            replace_with_equation(paragraph, math_elements[item[1]], item[1])
        else:
            paragraph.text = item[1]

    saved_path = DOCX_PATH
    try:
        doc.save(DOCX_PATH)
    except PermissionError:
        FALLBACK_PATH.parent.mkdir(parents=True, exist_ok=True)
        doc.save(FALLBACK_PATH)
        saved_path = FALLBACK_PATH

    with zipfile.ZipFile(saved_path) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Invalid docx zip member: {bad}")

    print(f"Updated methodology in {saved_path}")
    print(f"Backup: {backup}")
    print(f"Methodology slots replaced: {len(METHODOLOGY_ITEMS)}")
    print(f"OMML equations inserted: {len(FORMULAS)}")
    if saved_path != DOCX_PATH:
        print("Original was locked; fallback copy was written.")


if __name__ == "__main__":
    main()
