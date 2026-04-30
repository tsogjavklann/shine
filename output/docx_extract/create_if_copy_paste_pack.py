from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "bagiin_ner_v29.docx"
OUT = ROOT / "output" / "bagiin_ner_v29_Instrumental_Forest_copy_paste_pack.docx"
TABLE_DIR = ROOT / "output" / "tables"
FIG_DIR = ROOT / "output" / "figures"


def read_csv_dict(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


summary = {r["item"]: float(r["value"]) for r in read_csv_dict(TABLE_DIR / "T16_instrumental_forest_summary.csv")}
quartiles = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_tau_quartiles.csv")
regimes = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_threshold_regimes.csv")
importance = read_csv_dict(TABLE_DIR / "T16_instrumental_forest_variable_importance.csv")


def pct(x: float) -> str:
    return f"{float(x):.2f}"


def num(x: float, digits: int = 2) -> str:
    return f"{float(x):.{digits}f}"


def exp_return(beta: float) -> float:
    return 100 * (2.718281828459045 ** float(beta) - 1)


def clear_document(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def set_page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)


def set_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def run_font(run, size: int | None = None, bold: bool | None = None, italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc: Document, text: str = "", align=None, bold=False, italic=False, size=None, first_line=True, color=None):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    if first_line and text:
        p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(text)
    run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    run_font(r, size=14 if level == 1 else 12, bold=True)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_instruction(doc: Document, text: str) -> None:
    p = add_para(doc, text, first_line=False, bold=True, size=10, color="7A3E00")
    shade_paragraph(p, "FFF2CC")


def add_copy_label(doc: Document, text: str) -> None:
    p = add_para(doc, text, first_line=False, bold=True, size=10, color="1F4E79")
    shade_paragraph(p, "D9EAF7")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    run_font(r, size=10, bold=True, italic=True)


def add_source(doc: Document, text: str = "Эх сурвалж: Оюутны тооцоолол.") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    run_font(r, size=9, italic=True)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade_cell(hdr[i], "D9EAF7")
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                run_font(r, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    run_font(r, size=10)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    return table


def add_picture(doc: Document, image: Path, width_inches: float = 6.25) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image), width=Inches(width_inches))


def page_break(doc: Document) -> None:
    doc.add_page_break()


doc = Document(str(TEMPLATE))
clear_document(doc)
set_styles(doc)
set_page_setup(doc)

add_heading(doc, "Instrumental Forest нэмэлт шинжилгээ: үндсэн тайланд шууд paste хийх бэлэн хэсгүүд", 1)
add_instruction(
    doc,
    "Энэ файл нь standalone тайлан биш. Доорх БЛОК бүрийн цэнхэр 'COPY ЭХЛЭХ' шошгоны дараах агуулгыг үндсэн bagiin_ner_v29.docx-ийн заасан байрлалд шууд copy-paste хийхээр бэлтгэсэн. Шар зааврыг үндсэн тайланд бүү хуул.",
)
add_table(
    doc,
    ["Блок", "Үндсэн тайланд оруулах байрлал", "Юу нэмэх вэ"],
    [
        ["A", "ХУРААНГУЙ хэсэгт, одоогийн 2-р догол мөрийн дараа", "Instrumental Forest-ийн 1 богино үр дүн"],
        ["B", "ОРШИЛ -> Судалгааны зорилго, зорилтууд / Судалгааны шинэлэг тал", "Зорилт ба шинэлэг талын нэмэлт өгүүлбэрүүд"],
        ["C", "II БҮЛЭГ -> 2.5-ийн дараа, III БҮЛЭГ эхлэхээс өмнө", "2.6 Instrumental Forest аргын арга зүй"],
        ["D", "IV БҮЛЭГ -> 4.6-ийн төгсгөлд, V БҮЛЭГ эхлэхээс өмнө", "4.7 үр дүн, 3 хүснэгт, 3 зураг"],
        ["E", "V БҮЛЭГ -> 5.1 болон 5.3 хэсэгт", "Дүгнэлт ба хязгаарлалтын нэмэлт догол"],
        ["F", "НОМ ЗҮЙ хэсгийн төгсгөлд", "Нэмэх эх сурвалжууд"],
        ["G", "ХАВСРАЛТ -> Хавсралт В-ийн дараа", "Давтан ажиллуулах код ба output файлууд"],
    ],
    [1.3, 7.2, 6.8],
)
add_source(doc, "Тайлбар: Байрлалын зааврыг үндсэн тайлангийн одоогийн бүтэц, хүснэгт/зургийн дугаарлалттай тааруулан бэлтгэв.")
page_break(doc)

# Block A
add_heading(doc, "БЛОК A. Хураангуйд нэмэх догол мөр", 1)
add_instruction(doc, "Оруулах байрлал: ХУРААНГУЙ хэсгийн одоогийн хоёр дахь догол мөрийн дараа, 'Түлхүүр үгс'-ийн өмнө paste хийнэ.")
add_copy_label(doc, "COPY ЭХЛЭХ")
add_para(
    doc,
    f"Нэмэлтээр боловсролын өгөөжийн ялгаатай байдлыг нэг урьдчилан сонгосон босгоор хязгаарлахгүйгээр шалгах зорилгоор {int(summary['num_trees']):,} модтой Instrumental Forest шинжилгээ хийв. Уг шинжилгээгээр local IV өгөөж дунджаар {pct(summary['mean_return_pct'])} хувь, медианаар {pct(summary['median_return_pct'])} хувь, 10-90 хувийн мужид {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байна. Энэ нь ХШХБК-ийн дундаж өгөөжтэй ойролцоо боловч боловсролын өгөөж хувь хүний шинж, сургуулийн хүртээмж болон сурагч-багшийн харьцаатай хавсарсан олон хэмжээст ялгаатай байдлаар илэрч болохыг харуулж байна.",
)
add_copy_label(doc, "COPY ДУУСАХ")
page_break(doc)

# Block B
add_heading(doc, "БЛОК B. Оршилд нэмэх өгүүлбэрүүд", 1)
add_instruction(doc, "B1 байрлал: ОРШИЛ -> 'Судалгааны зорилго, зорилтууд' хэсэгт, одоогийн 'Дөрөвдүгээрт...' өгүүлбэрийн дараа нэмнэ.")
add_copy_label(doc, "B1 COPY ЭХЛЭХ")
add_para(
    doc,
    "Тавдугаарт, боловсролын өгөөжийн ялгаатай байдал нэг босгоны сонголтоос давж, хувь хүний болон сургуулийн орчны олон шинжтэй хамт илэрч байгаа эсэхийг Instrumental Forest аргаар нэмэлтээр зураглана.",
)
add_copy_label(doc, "B1 COPY ДУУСАХ")
add_instruction(doc, "B2 байрлал: ОРШИЛ -> 'Судалгааны шинэлэг тал' хэсгийн одоогийн догол мөрийн төгсгөлд нэмнэ.")
add_copy_label(doc, "B2 COPY ЭХЛЭХ")
add_para(
    doc,
    "Дөрөвдүгээрт, үндсэн босготой регрессийн үр дүнг машин сургалтын Instrumental Forest аргаар нэмэлтээр шалгаж, боловсролын өгөөжийн ялгаатай байдал зөвхөн нэг босгоны огтлол бус, хувь хүний шинж, сургуулийн хүртээмж, сурагч-багшийн харьцаа, төрсөн үе болон судалгааны давалгаатай хамт илэрч байгаа эсэхийг өгөгдөлд суурилсан байдлаар дүрсэлж байгаа нь судалгааны эрэл хайгуулын шинжийг нэмэгдүүлж байна.",
)
add_copy_label(doc, "B2 COPY ДУУСАХ")
add_instruction(doc, "B3 байрлал: ОРШИЛ -> 'Ажлын бүтэц' хэсгийн одоогийн догол мөрийн 'Дөрөвдүгээр бүлэгт...' өгүүлбэрийг солих эсвэл дараах өгүүлбэрээр баяжуулна.")
add_copy_label(doc, "B3 COPY ЭХЛЭХ")
add_para(
    doc,
    "Дөрөвдүгээр бүлэгт ЭХБК, ХШХБК, ХХБР үнэлгээ, бүүтстрап шалгалт болон Instrumental Forest-ийн нэмэлт эрэл хайгуулын үр дүнг нэгтгэн шинжилнэ.",
)
add_copy_label(doc, "B3 COPY ДУУСАХ")
page_break(doc)

# Block C
add_heading(doc, "БЛОК C. II бүлэгт нэмэх арга зүйн дэд хэсэг", 1)
add_instruction(doc, "Оруулах байрлал: II БҮЛЭГ-ийн 2.5 'Үр дүнгийн бат бөх байдлын шинжилгээний бүтэц' хэсгийн дараа, III БҮЛЭГ эхлэхийн өмнө paste хийнэ. Гарчгийн дугаарлалт үндсэн тайланд 2.6 болно.")
add_copy_label(doc, "COPY ЭХЛЭХ")
add_heading(doc, "2.6. Instrumental Forest нэмэлт эрэл хайгуулын арга", 2)
add_para(
    doc,
    "Үндсэн ХХБР загвар боловсролын өгөөж 17-18 насны сурагч-багшийн харьцааны тодорхой босгоор ялгаатай эсэхийг шалгадаг. Гэвч боловсролын өгөөжийн ялгаатай байдал зөвхөн нэг босгоны огтлолоор бүрэн тайлбарлагдахгүй байж болох тул нэмэлтээр Instrumental Forest буюу хэрэгсэл хувьсагчтай ой загварыг ашиглав. Энэ арга нь хэрэгсэл хувьсагчийн логикийг machine learning-ийн модон хуваалтын аргатай хослуулж, ажиглагдсан шинжүүдийн нөхцөлд боловсролын local IV өгөөж хэрхэн ялгаатай байгааг зураглах боломж олгодог.",
)
add_para(
    doc,
    "Instrumental Forest шинжилгээнд хамааруулах хувьсагч нь логаритмчилсан цалин, эндоген тайлбарлагч нь боловсролын жил, хэрэгсэл хувьсагч нь эцэг, эхийн боловсролын дундаж хэвээр байна. Харин heterogeneity-г ялгах шинжүүдэд нас, насны квадрат, хүйс, гэрлэлтийн байдал, хотын суурьшил, 17-18 насны сурагч-багшийн дундаж харьцаа, сургуулийн хүртээмж, төрсөн аймаг, төрсөн үеийн бүлэг болон судалгааны давалгааг оруулсан. Эцэг, эхийн боловсролын дундаж нь хэрэгсэл хувьсагчийн үүрэгтэй тул X шинжийн багцад давхар оруулаагүй.",
)
add_para(
    doc,
    "Энэхүү нэмэлт шинжилгээг үндсэн шалтгаант нотолгоо гэж бус, эрэл хайгуулын heterogeneity map гэж тайлбарлана. Өөрөөр хэлбэл, Instrumental Forest-ийн гаргасан өгөөж нь тухайн хүний яг хувийн шалтгаант нөлөө бус, ижил төстэй шинж бүхий бүлгийн нөхцөлт IV өгөөжийн таамаглал юм. Иймээс уг арга нь ХХБР загварыг орлохгүй, харин боловсролын өгөөжийн ялгаа сургуулийн орчин болон хувь хүний шинжүүдтэй хэрхэн хавсарч байгааг нэмэлтээр шалгах үүрэгтэй.",
)
add_para(
    doc,
    f"Загварыг R орчны grf багцын instrumental_forest функцээр {int(summary['num_trees']):,} модтойгоор үнэлэв. Үр дүнгийн тайлбар нь гурван хэсгээс бүрдэнэ: local IV өгөөжийн тархалт, сурагч-багшийн харьцаатай холбогдох дүр зураг, өгөөжийн ялгааг тайлбарлах хувьсагчдын ач холбогдлын нэгтгэл.",
)
add_copy_label(doc, "COPY ДУУСАХ")
page_break(doc)

# Block D
add_heading(doc, "БЛОК D. IV бүлэгт нэмэх үр дүнгийн хэсэг", 1)
add_instruction(doc, "Оруулах байрлал: IV БҮЛЭГ-ийн 4.6 'Эконометрик үр дүнгийн нэгдсэн дүгнэлт' хэсгийн төгсгөлд, V БҮЛЭГ эхлэхээс өмнө paste хийнэ. Одоогийн тайланд хүснэгт 8, зураг 5 хүртэл байгаа тул доорх дугаарлалтыг шууд үргэлжлүүлсэн.")
add_copy_label(doc, "COPY ЭХЛЭХ")
add_heading(doc, "4.7. Instrumental Forest нэмэлт эрэл хайгуулын шинжилгээ", 2)
add_para(
    doc,
    "Үндсэн ХХБР үнэлгээ боловсролын өгөөжийг 17-18 насны сурагч-багшийн дундаж харьцааны босгоор хоёр регимд хуваан харуулсан. Харин боловсролын өгөөжийн ялгаа зөвхөн нэг босгоор бус, хувь хүний нас, хүйс, гэрлэлтийн байдал, хотжилт, сургуулийн хүртээмж, төрсөн үе, судалгааны давалгаа зэрэг олон шинжтэй хамт илэрч болох тул Instrumental Forest аргаар нэмэлт эрэл хайгуулын шинжилгээ хийв. Энэ шинжилгээ үндсэн ХХБР үр дүнг орлохгүй, харин боловсролын өгөөжийн ялгаа өгөгдөл дотор хэрхэн тархаж байгааг дүрслэх зорилготой.",
)
add_caption(doc, "Хүснэгт 9. Instrumental Forest local IV өгөөжийн үндсэн статистик")
add_table(
    doc,
    ["Үзүүлэлт", "Утга", "Тайлбар"],
    [
        ["Ажиглалтын тоо", f"{int(summary['N']):,}", "ХХБР эцсийн түүвэртэй ижил"],
        ["Модны тоо", f"{int(summary['num_trees']):,}", "grf::instrumental_forest"],
        ["X шинжийн тоо", f"{int(summary['X_features'])}", "Хувь хүн, сургууль, үе, давалгааны шинжүүд"],
        ["Дундаж local IV өгөөж", f"{pct(summary['mean_return_pct'])}%", "Forest-ийн дундаж таамаглал"],
        ["Медиан local IV өгөөж", f"{pct(summary['median_return_pct'])}%", "Тархалтын төв"],
        ["10-90 хувийн муж", f"{pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])}%", "Өгөөжийн ялгаатай байдлын хүрээ"],
        ["Харьцуулах ХШХБК өгөөж", f"{pct(exp_return(summary['IV_2SLS_beta_log_reference']))}%", "Үндсэн IV дундаж өгөөж"],
    ],
    [4.6, 3.2, 7.0],
)
add_source(doc)
add_para(
    doc,
    f"Хүснэгт 9-өөс харахад Instrumental Forest-ийн таамагласан local IV өгөөж дунджаар {pct(summary['mean_return_pct'])} хувь, медианаар {pct(summary['median_return_pct'])} хувь байна. Энэ нь ХШХБК үнэлгээнээс гарсан {pct(exp_return(summary['IV_2SLS_beta_log_reference']))} хувийн дундаж өгөөжтэй ойролцоо байгаа тул нэмэлт machine learning загвар нийт түвшинд үндсэн IV үр дүнгээс зөрсөн, хэт өөр дүгнэлт гаргаагүй байна. Харин 10-90 хувийн муж {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байгаа нь боловсролын өгөөж нэг дундаж коэффициентээр бүрэн тайлбарлагдахгүй, бүлэг хоорондын ялгаатай байдлыг агуулж байгааг харуулна.",
)
add_caption(doc, "Зураг 6. Instrumental Forest local IV өгөөжийн тархалт")
add_picture(doc, FIG_DIR / "instrumental_forest_tau_distribution.png")
add_source(doc)
add_para(
    doc,
    "Зураг 6-д local IV өгөөжийн тархалтыг харуулав. Тасархай босоо шугам нь ХШХБК-ийн дундаж өгөөжийг илэрхийлнэ. Тархалтын төв хэсэг 10-12 хувийн орчимд байрлаж байгаа боловч баруун талдаа өндөр өгөөжтэй бүлгүүд ажиглагдаж байна. Энэ нь боловсролын өгөөжийн дундаж хэмжээ эерэг бөгөөд бодлогын хувьд чухал боловч бүх ажиллагчдад ижил хүчтэй илэрдэггүйг харуулна.",
)
add_caption(doc, "Хүснэгт 10. Local IV өгөөжийн квартил бүлгүүдийн харьцуулалт")
add_table(
    doc,
    ["Бүлэг", "N", "Дундаж өгөөж", "STR медиан", "Хотын хувь", "Эмэгтэй хувь"],
    [
        [
            r["tau_group_label"],
            f"{int(float(r['N'])):,}",
            f"{pct(float(r['mean_return_pct']))}%",
            num(float(r["median_student_teacher_ratio_17_18"])),
            f"{pct(100 * float(r['urban_share']))}%",
            f"{pct(100 * float(r['female_share']))}%",
        ]
        for r in quartiles
    ],
    [3.0, 1.7, 2.6, 2.5, 2.3, 2.3],
)
add_source(doc)
add_para(
    doc,
    "Хүснэгт 10-д local IV өгөөжийн таамаглалаар дөрвөн тэнцүү бүлэг болгон харьцуулсан үр дүнг үзүүлэв. Доод 25 хувийн бүлгийн дундаж өгөөж 6.95 хувь байгаа бол дээд 25 хувийн бүлгийн дундаж өгөөж 16.93 хувь байна. Энэ ялгаа нь Instrumental Forest-ийн гол мэдээлэл бөгөөд боловсролын өгөөж нь зөвхөн дундаж коэффициент бус, ажиглагдсан шинжүүдээр ялгаатай тархах боломжтойг харуулна.",
)
add_caption(doc, "Зураг 7. Local IV өгөөж ба 17-18 насны сурагч-багшийн харьцаа")
add_picture(doc, FIG_DIR / "instrumental_forest_tau_by_student_teacher_ratio.png")
add_source(doc)
add_para(
    doc,
    "Зураг 7-д local IV өгөөжийг 17-18 насны сурагч-багшийн дундаж харьцаатай холбон харуулав. Босгоны орчимд ажиглалтууд нягт байрлаж байгаа бөгөөд өгөөж нь нэг чиглэлтэй шугаман өсөлт эсвэл бууралтаар бүрэн тайлбарлагдахгүй байна. STR-ийн 19.53 босгоор хоёр регимд дундажлахад local IV өгөөж доод регимд 11.99 хувь, дээд регимд 11.04 хувь гарсан. Энэ нь ХХБР-ийн дээд регимд өндөр өгөөж илэрсэн үр дүнг шууд давтаагүй боловч боловсролын өгөөжийн ялгаа зөвхөн STR босгоор бус, бусад шинжүүдтэй хамт нөхцөлдөн илэрч байгааг харуулна.",
)
add_caption(doc, "Хүснэгт 11. Instrumental Forest variable importance-ийн нэгтгэл")
add_table(
    doc,
    ["Хувьсагчийн бүлэг", "Харьцангуй ач холбогдол", "Тайлбар"],
    [
        [r["family"], f"{pct(100 * float(r['share']))}%", "Өгөөжийн ялгааг ялгахад ашиглагдсан мэдээллийн хувь"]
        for r in importance
    ],
    [5.0, 3.4, 7.0],
)
add_source(doc)
add_caption(doc, "Зураг 8. Өгөөжийн ялгааг тайлбарлах хувьсагчдын ач холбогдол")
add_picture(doc, FIG_DIR / "instrumental_forest_variable_importance.png", width_inches=6.0)
add_source(doc)
add_para(
    doc,
    "Variable importance-ийн үр дүнгээр өгөөжийн ялгааг ялгахад хувийн шинжүүд 48.02 хувь, сургуулийн хүртээмж 19.23 хувь, сурагч-багшийн харьцаа 15.23 хувь, судалгааны давалгаа 10.22 хувь, төрсөн үе 7.29 хувийн харьцангуй ач холбогдолтой байна. Энэ нь сурагч-багшийн харьцаа боловсролын өгөөжийн ялгааг тайлбарлахад мэдээлэл өгч байгаа боловч ганцаараа бүх heterogeneity-г тодорхойлохгүйг харуулна. Иймээс ХХБР загварын босго үр дүнг боловсролын өгөөжийн нэг боломжит огтлол гэж үзэхийн зэрэгцээ Instrumental Forest-ийн үр дүнгээр уг ялгаа олон хэмжээст шинжтэй болохыг нэмэлтээр баталгаажуулж байна.",
)
add_para(
    doc,
    "Нэгтгэн дүгнэвэл, Instrumental Forest-ийн нэмэлт шинжилгээ нь боловсролын өгөөжийн дундаж таамаглал үндсэн ХШХБК үнэлгээтэй ойролцоо боловч өгөөж хувь хүн, сургуулийн хүртээмж, сурагч-багшийн харьцаа, судалгааны давалгаа болон төрсөн үеийн шинжүүдээр ялгаатай тархаж болохыг харууллаа. Энэ үр дүн нь ХХБР загварын шалтгаант дүгнэлтийг орлохгүй, харин боловсролын өгөөжийн ялгаатай байдал нэг босгоос давсан олон хүчин зүйлийн нийлбэр байж болохыг дэмжсэн эрэл хайгуулын нотолгоо юм.",
)
add_copy_label(doc, "COPY ДУУСАХ")
page_break(doc)

# Block E
add_heading(doc, "БЛОК E. Дүгнэлт ба хязгаарлалтад нэмэх догол мөрүүд", 1)
add_instruction(doc, "E1 байрлал: V БҮЛЭГ -> 5.1 'Үндсэн дүгнэлт' хэсэгт, одоогийн 'Гуравдугаарт...' догол мөрийн дараа нэмнэ.")
add_copy_label(doc, "E1 COPY ЭХЛЭХ")
add_para(
    doc,
    f"Дөрөвдүгээрт, Instrumental Forest-ийн нэмэлт шинжилгээгээр боловсролын local IV өгөөж дунджаар {pct(summary['mean_return_pct'])} хувь, медианаар {pct(summary['median_return_pct'])} хувь гарч, үндсэн ХШХБК үнэлгээтэй ойролцоо байна. Гэхдээ өгөөжийн 10-90 хувийн муж {pct(summary['p10_return_pct'])}-{pct(summary['p90_return_pct'])} хувь байгаа нь боловсролын өгөөж нэг дундаж коэффициентээр бүрэн тайлбарлагдахгүй, хувь хүний шинж, сургуулийн хүртээмж, сурагч-багшийн харьцаа, судалгааны давалгаа болон төрсөн үеийн ялгаатай хамт илэрч болохыг харууллаа.",
)
add_copy_label(doc, "E1 COPY ДУУСАХ")
add_instruction(doc, "E2 байрлал: V БҮЛЭГ -> 5.3 'Судалгааны хязгаарлалт...' хэсгийн төгсгөлд, дүгнэж хэлбэл өгүүлбэрийн өмнө нэмнэ.")
add_copy_label(doc, "E2 COPY ЭХЛЭХ")
add_para(
    doc,
    "Дөрөвдүгээрт, Instrumental Forest-ийн үр дүнг статистикийн уламжлалт нэг p-утгатай баталгаа гэж тайлбарлахгүй. Уг арга нь өгөөжийн ялгаатай байдлыг өгөгдөлд суурилсан хэлбэрээр зураглах нэмэлт machine learning шинжилгээ бөгөөд эцэг, эхийн боловсролын дундаж хэрэгсэл хувьсагчийн хязгаарлалтын таамаглал хэвээр хадгалагдана. Иймээс Instrumental Forest-ийн үр дүнг үндсэн ХХБР-ийн шалтгаант тайлбарыг орлох бус, боловсролын өгөөжийн heterogeneity олон хэмжээст шинжтэй байж болохыг дэмжсэн exploratory evidence гэж ашиглана.",
)
add_copy_label(doc, "E2 COPY ДУУСАХ")
page_break(doc)

# Block F
add_heading(doc, "БЛОК F. Ном зүйд нэмэх эх сурвалж", 1)
add_instruction(doc, "Оруулах байрлал: НОМ ЗҮЙ хэсгийн төгсгөлд, одоогийн эх сурвалжуудын дараа paste хийнэ. Дугаарлалтыг үндсэн ном зүйн дарааллаар үргэлжлүүлж болно.")
add_copy_label(doc, "COPY ЭХЛЭХ")
refs = [
    "Athey, S., Tibshirani, J., & Wager, S. (2019). Generalized random forests. Annals of Statistics, 47(2), 1148-1178.",
    "Chernozhukov, V., Chetverikov, D., Demirer, M., Duflo, E., Hansen, C., Newey, W., & Robins, J. (2018). Double/debiased machine learning for treatment and structural parameters. Econometrics Journal, 21(1), C1-C68.",
    "Tibshirani, J., Athey, S., Friedberg, R., Hadad, V., Hirshberg, D., Miner, L., Sverdrup, E., Wager, S., Wright, M., & grf contributors. (2024). grf: Generalized Random Forests. R package.",
]
for ref in refs:
    add_para(doc, ref, first_line=False)
add_copy_label(doc, "COPY ДУУСАХ")
page_break(doc)

# Block G
add_heading(doc, "БЛОК G. Хавсралтад нэмэх кодын тайлбар", 1)
add_instruction(doc, "Оруулах байрлал: ХАВСРАЛТ -> Хавсралт В. Программ хангамжийн код хэсгийн дараа paste хийнэ.")
add_copy_label(doc, "COPY ЭХЛЭХ")
add_heading(doc, "Хавсралт Г. Instrumental Forest нэмэлт шинжилгээний код ба гаралт", 2)
add_para(
    doc,
    "Instrumental Forest нэмэлт шинжилгээг R/31_instrumental_forest.R script-ээр гүйцэтгэв. Уг script нь эцсийн ХХБР түүвэр болох data/processed/ivtr_ready_parent_educ_mean_student_teacher_avg_17_18.rds файлыг уншиж, grf багцын instrumental_forest функцээр боловсролын local IV өгөөжийг тооцсон. Үр дүнгийн хүснэгтүүд output/tables хавтаст, зургууд output/figures хавтаст хадгалагдсан.",
)
add_para(doc, "Давтан ажиллуулах команд: $env:GRF_TREES='2000'; $env:GRF_THREADS='3'; Rscript R/31_instrumental_forest.R", first_line=False)
add_para(
    doc,
    "Гол гаралтын файлууд: T16_instrumental_forest_summary.csv, T16_instrumental_forest_tau_quartiles.csv, T16_instrumental_forest_threshold_regimes.csv, T16_instrumental_forest_variable_importance.csv, instrumental_forest_tau_distribution.png, instrumental_forest_tau_by_student_teacher_ratio.png, instrumental_forest_variable_importance.png.",
)
add_copy_label(doc, "COPY ДУУСАХ")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
